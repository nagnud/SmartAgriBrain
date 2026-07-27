from __future__ import annotations

import shutil
import sys
from copy import deepcopy
from pathlib import Path
import re
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree


BIBLIOGRAPHY_NS = "http://schemas.openxmlformats.org/officeDocument/2006/bibliography"


def replace_paragraph_text(paragraph, text: str) -> None:
    """Replace paragraph text while preserving its paragraph and run formatting."""
    first_run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        first_run_properties = deepcopy(paragraph.runs[0]._r.rPr)

    for child in list(paragraph._p):
        if child.tag.endswith("}pPr"):
            continue
        paragraph._p.remove(child)

    run = paragraph.add_run(text)
    if first_run_properties is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, first_run_properties)


def replace_toc_hyperlink_text(paragraph, text: str) -> None:
    """Replace one static-TOC entry without changing its bookmark target."""
    hyperlink = paragraph._p.find(qn("w:hyperlink"))
    if hyperlink is None:
        raise RuntimeError(f"TOC paragraph has no hyperlink: {paragraph.text}")

    first_run = hyperlink.find(qn("w:r"))
    run_properties = deepcopy(first_run.find(qn("w:rPr"))) if first_run is not None else None
    for child in list(hyperlink):
        hyperlink.remove(child)

    run = OxmlElement("w:r")
    if run_properties is not None:
        run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)


def replace_bookmarked_heading_text(paragraph, text: str) -> None:
    """Replace heading text while preserving the bookmark used by the TOC."""
    first_run = paragraph._p.find(qn("w:r"))
    run_properties = deepcopy(first_run.find(qn("w:rPr"))) if first_run is not None else None
    for child in list(paragraph._p):
        if child.tag in {qn("w:pPr"), qn("w:bookmarkStart"), qn("w:bookmarkEnd")}:
            continue
        paragraph._p.remove(child)

    run = OxmlElement("w:r")
    if run_properties is not None:
        run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)

    bookmark_end = paragraph._p.find(qn("w:bookmarkEnd"))
    if bookmark_end is None:
        paragraph._p.append(run)
    else:
        paragraph._p.insert(paragraph._p.index(bookmark_end), run)


def add_internal_citation_links(paragraph) -> None:
    """Turn every [n] citation marker into a hyperlink to bookmark ref_n."""
    text = paragraph.text
    matches = list(re.finditer(r"\[([1-9][0-9]*)\]", text))
    if not matches:
        return

    first_run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        first_run_properties = deepcopy(paragraph.runs[0]._r.rPr)
    for child in list(paragraph._p):
        if child.tag.endswith("}pPr"):
            continue
        paragraph._p.remove(child)

    def append_run(value: str, hyperlink=None, hyperlink_style: bool = False) -> None:
        if not value:
            return
        run = OxmlElement("w:r")
        if first_run_properties is not None:
            run.append(deepcopy(first_run_properties))
        if hyperlink_style:
            run_properties = run.find(qn("w:rPr"))
            if run_properties is None:
                run_properties = OxmlElement("w:rPr")
                run.insert(0, run_properties)
            run_style = OxmlElement("w:rStyle")
            run_style.set(qn("w:val"), "Hyperlink")
            run_properties.insert(0, run_style)
        text_element = OxmlElement("w:t")
        text_element.text = value
        run.append(text_element)
        (hyperlink if hyperlink is not None else paragraph._p).append(run)

    cursor = 0
    for match in matches:
        append_run(text[cursor:match.start()])
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("w:anchor"), f"ref_{match.group(1)}")
        hyperlink.set(qn("w:history"), "1")
        append_run(match.group(0), hyperlink=hyperlink, hyperlink_style=True)
        paragraph._p.append(hyperlink)
        cursor = match.end()
    append_run(text[cursor:])


def ensure_paragraph_bookmark(document, paragraph, bookmark_name: str) -> None:
    """Wrap a reference paragraph in a bookmark used by citation hyperlinks."""
    existing = paragraph._p.xpath(f'./w:bookmarkStart[@w:name="{bookmark_name}"]')
    if existing:
        return

    bookmark_ids = []
    for element in document._element.body.xpath(".//w:bookmarkStart"):
        raw_id = element.get(qn("w:id"))
        if raw_id is not None and raw_id.isdigit():
            bookmark_ids.append(int(raw_id))
    bookmark_id = str(max(bookmark_ids, default=-1) + 1)

    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bookmark_id)
    start.set(qn("w:name"), bookmark_name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bookmark_id)

    insert_index = 1 if paragraph._p.find(qn("w:pPr")) is not None else 0
    paragraph._p.insert(insert_index, start)
    paragraph._p.append(end)


def patch_bibliography(docx_path: Path) -> None:
    """Replace the obsolete DS18B20 Word bibliography source with DHT11."""
    temporary_path = docx_path.with_suffix(".bibliography.tmp.docx")
    namespace = {"b": BIBLIOGRAPHY_NS}

    with ZipFile(docx_path, "r") as source_archive:
        bibliography_xml = source_archive.read("customXml/item1.xml")
        root = etree.fromstring(bibliography_xml)

        target_source = None
        for source in root.findall("b:Source", namespace):
            tag = source.findtext("b:Tag", namespaces=namespace)
            if tag in {"ADI2025DS18B20", "ASAIRDHT11"}:
                target_source = source
                break
        if target_source is None:
            raise RuntimeError("Neither the DS18B20 nor DHT11 bibliography source was found.")

        replacements = {
            "Tag": "ASAIRDHT11",
            "Guid": "{5F3E2636-6D44-4EAF-98F7-4B68835824AC}",
            "Title": "DHT11 Temperature and Humidity Sensor Module",
            "Year": "2026",
            "Publisher": "ASAIR",
            "URL": "https://www.aosong.com/en/Products/info.aspx?itemid=2257",
            "YearAccessed": "2026",
            "MonthAccessed": "7",
            "DayAccessed": "25",
        }
        for field_name, value in replacements.items():
            element = target_source.find(f"b:{field_name}", namespace)
            if element is None:
                element = etree.SubElement(target_source, f"{{{BIBLIOGRAPHY_NS}}}{field_name}")
            element.text = value

        corporate = target_source.find(".//b:Corporate", namespace)
        if corporate is None:
            raise RuntimeError("The bibliography source has no corporate-author node.")
        corporate.text = "ASAIR"

        updated_bibliography_xml = etree.tostring(
            root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        with ZipFile(temporary_path, "w", compression=ZIP_DEFLATED) as target_archive:
            for item in source_archive.infolist():
                if item.filename == "customXml/item1.xml":
                    target_archive.writestr(item, updated_bibliography_xml)
                else:
                    target_archive.writestr(item, source_archive.read(item.filename))

    temporary_path.replace(docx_path)


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: update_design_doc.py INPUT.docx BACKUP.docx")

    input_path = Path(sys.argv[1]).resolve()
    backup_path = Path(sys.argv[2]).resolve()
    backup_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(input_path, backup_path)

    document = Document(input_path)
    replacements = {
        12: (
            "面向家庭阳台、室内种植空间和小型庭院中养护信息分散、人工观察不连续、"
            "浇灌位置难以准确控制等问题，本文设计并实现SmartAgriBrain智慧农业系统。"
            "产品采用“小型智慧种植屋”式方形一体化结构，将DHT11温湿度传感器、BH1750"
            "光照传感器、二氧化碳传感器、补光模块、二维舵机云台、定点灌溉水枪、"
            "ESP32-C5语音显示终端与普通ESP32采集执行节点集成于同一设备。普通ESP32"
            "持续采集空气温度、空气相对湿度、光照强度和二氧化碳浓度，并通过Wi-Fi、"
            "TLS和云端EMQX以MQTT QoS 1上报；系统可依据用户设置和养护策略调节补光亮度，"
            "按照目标距离与方位控制二维云台和水泵完成定点浇灌。显示屏用于呈现环境参数、"
            "设备状态与养护信息，Web端提供实时监测、历史查询、远程控制、图像分析和知识库"
            "管理，ESP32-C5通过云端大模型完成语音问答，为植物识别、异常分析、浇水施肥及"
            "日常养护提供建议。系统由此形成集环境监测、精准执行、可视化展示、远程管理、"
            "语音交互和智能决策于一体的家庭智能植物管家。"
        ),
        13: "关键词：家庭智能种植；ESP32；MQTT；精准灌溉；语音交互",
        62: (
            "本产品定位为面向家庭阳台、室内种植空间及小型庭院的智能植物养护终端，"
            "采用方形一体化“小型智慧种植屋”设计。与面向大规模温室的生产管理平台不同，"
            "本系统强调占地紧凑、部署简单、状态直观和家庭用户易操作，将环境感知、补光、"
            "定点浇灌、屏幕显示、Web远程管理及语音问答集中在同一设备中，使用户在缺少"
            "专业种植经验时也能持续了解植物状态并完成日常养护。"
        ),
        64: (
            "现场传感器采用不同的接口和数据单位：DHT11通过GPIO4提供空气温度与相对湿度，"
            "BH1750通过I2C提供光照强度，JW01模块通过串口帧提供二氧化碳浓度。系统在MQTT"
            "与REST边界统一使用temperature_c、humidity_pct、illuminance_lux和co2_ppm，"
            "并为每项数据附加质量状态。固定字段、单位、时间戳和无效值语义后，显示屏、"
            "Web历史曲线、告警分析和大模型上下文可以直接复用同一份领域数据。"
        ),
        67: (
            "家庭网络可能短时波动，设备主循环不能因一次传感器读取或网络重连长期阻塞。"
            "普通ESP32通过非阻塞调度持续维护Wi-Fi、MQTT、水枪状态机和传感器采样；MQTT"
            "采用发布/订阅模型，使现场设备、云端服务、Web和C5终端能够异步协作[1]。"
            "控制消息采用QoS 1、有效期、命令标识和设备确认机制，网络恢复后也不会把发布成功"
            "错误理解为硬件已经执行。"
        ),
        69: (
            "水泵、补光灯和二维云台均直接作用于实体硬件，必须限制输出范围并设置失效策略。"
            "GPIO26水泵和GPIO14补光灯经驱动模块使用PWM控制；GPIO27水平SG90限制在0至180度，"
            "GPIO13俯仰SG90依据机构约束限制在5至60度，两个舵机使用独立稳定5 V电源并与ESP32"
            "共地。目标变化时系统先停泵、再转向、等待稳定后开泵；网络断开、定时结束或动态"
            "会话超过3秒未保活时立即停泵，避免失联持续喷水。"
        ),
        71: "形成可持续采集空气温度、空气相对湿度、光照强度和二氧化碳浓度的家庭种植环境数据源。",
        72: "支持补光灯亮度调节，以及由目标距离、方位、静态或动态模式和喷水时长共同描述的定点浇灌控制。",
        73: "通过云端EMQX、统一MQTT契约、REST/SSE接口和设备ACK打通普通ESP32、后端、Web与C5终端。",
        74: "提供屏幕状态展示、Web实时与历史数据、植物知识库、图像分析和云端大模型语音养护问答。",
        78: (
            "系统采用双MCU分工。普通ESP32负责DHT11、BH1750和JW01等环境数据采集，并实时控制"
            "补光灯、水泵和二维云台；ESP32-C5作为设备正面的语音与显示终端，负责录音、播放、"
            "屏幕界面以及与云端大模型服务交互。硬实时的采集执行任务与交互任务相互隔离，既避免"
            "语音和显示占用影响水泵安全状态机，也便于独立升级传感器、屏幕和云端模型能力。"
        ),
        81: (
            "系统将遥测、控制命令和设备确认划分为独立消息。遥测描述环境与执行器快照；命令描述"
            "目标、期望值和有效期；确认只在设备完成校验并应用输出或明确拒绝后发布。MQTT PUBACK"
            "只表示Broker已接收消息，后端返回queued或dispatched也不代表硬件动作完成；Web只有收到"
            "与command_id对应的executed确认后才把操作视为成功。SG90和水路目前没有闭环反馈，"
            "因此确认表示控制输出已应用，不等同于测得真实角度或流量。"
        ),
        84: (
            "系统把两个SG90、水泵和喷嘴组合为一个目标喷水执行器。正式target_position命令同时携带"
            "ground_range_mm、bearing_deg、static或dynamic模式、泵控制百分比、会话序号与定时参数。"
            "设备按“停泵、转向、稳定等待、开泵”的顺序原子执行，防止云台和水路命令乱序。动态模式"
            "只在目标发生变化时重写舵机PWM；未喷水时Web心跳只维持后端会话，喷水时设备保活仅刷新"
            "安全超时，相同位置不会驱动舵机往复动作。"
        ),
        88: (
            "现场ESP32周期采集空气温度、空气相对湿度、光照强度和二氧化碳浓度。DHT11的数据线接"
            "GPIO4，驱动以不短于2秒的间隔一次获得温度与湿度，读取失败时上报null并将质量标记为"
            "invalid；BH1750通过I2C输出Lux；JW01模块通过串口帧解析ppm。BH1750的数字光照输出适合"
            "作为补光判断输入[2]，DHT11能够以单总线数字信号同时提供家庭种植空间的温湿度数据[3]。"
        ),
        91: (
            "正式控制Topic为smartagribrain/v1/devices/{device_id}/command。独立set命令仅用于"
            "grow_light，协议值0至90映射到GPIO14补光灯PWM。水泵和两个舵机是水枪内部部件，不接受"
            "彼此分离的远程命令；定点浇灌必须使用target_position一次携带位置、水泵、模式和时长，"
            "由GPIO26、GPIO27和GPIO13的统一安全状态机执行。浏览器不保存MQTT凭据，也不直接连接"
            "设备，所有真实控制均由后端校验后发布。"
        ),
        93: (
            "水平舵机使用GPIO27，Web方位角-90至90度线性对应舵机真实命令角0至180度，0度方位对应"
            "舵机90度正前方。俯仰舵机使用GPIO13，受实体机构限制只允许5至60度：5度表示水枪平行"
            "向前，60度表示竖直向上；目标地面距离300至1200毫米由近到远线性映射为60至5度。"
            "上电初始位置为水平90度、俯仰5度，SG90采用50 Hz PWM驱动[4]。该几何关系保证软件边界"
            "一致，最终落点仍需结合喷嘴安装、水压与弹道进行实物标定。"
        ),
        96: (
            "FastAPI后端是Web、普通ESP32和C5终端的统一服务入口，负责订阅设备遥测、保存SQLite"
            "历史数据和命令状态、通过云端EMQX发布真实控制，并提供REST与SSE接口。Vue Web端展示"
            "实时参数、历史趋势、在线状态、告警和执行器状态，提供补光与水枪控制、摄像头画面、"
            "植物识别、位置分析、知识库管理和大模型养护建议。前端不直接连接EMQX，敏感凭据只保存在"
            "后端和设备本机配置中。"
        ),
        99: (
            "ESP32-C5终端面向设备旁的直接交互，显示当前环境参数、执行器状态和养护信息，并将语音"
            "请求经统一云端服务送入大模型。后端已定义C5页面状态和助手响应Topic，并提供语音识别、"
            "对话编排、知识库检索与语音合成所需接口；终端只负责采集、播放和展示，不直接操作普通"
            "ESP32的GPIO。涉及浇灌等高风险动作时，模型输出必须转换为结构化候选命令并经过确认，"
            "不能由自然语言直接绕过控制校验。"
        ),
        103: (
            "系统由一体化设备层、网络传输层、云端服务层和用户交互层组成。设备层在方形“小型智慧"
            "种植屋”内集成传感器、补光灯、水枪、二维云台、普通ESP32以及C5语音显示终端；网络层使用"
            "Wi-Fi、TLS和MQTT；云端层负责EMQX接入、数据持久化、知识库、图像分析和模型服务；交互层"
            "包括设备屏幕、语音和Vue Web。各层通过明确的数据对象协作，使家庭用户获得一体化体验，"
            "同时保持传感器、执行器和模型服务可独立替换。"
        ),
        105: (
            "环境数据流为“DHT11/BH1750/JW01—普通ESP32—EMQX—FastAPI—数据库—Web或C5屏幕”；"
            "控制流为“Web或C5语音确认—FastAPI—EMQX—普通ESP32—补光灯或水枪—command_ack”；"
            "智能分析流为“实时数据、摄像头图像和植物知识库—云端模型—结构化建议—Web或C5”。"
            "三条链路共享站点与设备标识，但分别保留数据、控制和建议的语义边界。"
        ),
        108: (
            "普通ESP32中，GPIO4连接DHT11数据线，GPIO26用于水泵PWM，GPIO14用于补光灯PWM，GPIO27"
            "用于水平舵机，GPIO13用于俯仰舵机。水泵和补光灯采用1 kHz、8位LEDC通道；两个SG90分别"
            "采用50 Hz、16位LEDC通道。执行器通道互不复用，传感器读取与PWM更新均由主循环按非阻塞"
            "方式调度。Arduino-ESP32提供LEDC配置、引脚绑定和占空比写入接口[5]。"
        ),
        110: (
            "PanTilt类封装双SG90的50 Hz PWM初始化、角度边界、脉宽换算和当前命令角保存；"
            "WaterGunController进一步把云台和GPIO26水泵组合为非阻塞状态机。只有目标距离或方位"
            "发生变化时才写入两个舵机通道，相同目标的动态保活不会重复移动。控制器在800毫秒稳定"
            "等待后决定是否开泵，并独立处理定时截止、3秒动态超时、断网急停和最终ACK。"
        ),
        113: (
            "普通ESP32和FastAPI均通过TLS连接云端EMQX，不启动也不回退到本地Mosquitto。设备使用"
            "MQTT 3.1.1持久会话，command、command_ack、telemetry、status、capabilities及LWT均采用"
            "QoS 1；status和capabilities使用retain保存最新在线与能力信息，command禁止retain以避免"
            "重连误执行。设备端网络回调只负责收包入队，Arduino主循环完成JSON校验和硬件分发，"
            "MQTT基础交互遵循标准[1]。"
        ),
        115: (
            "mqtt_client模块使用ArduinoJson解析精简命令信封。Topic已确定协议版本和设备，因此载荷"
            "只保留command_id、expires_at和command；普通ESP32校验精确Topic、本地时间、有效期、"
            "参数范围、动态session_id与递增sequence，并以内容指纹和ACK缓存处理QoS 1重复投递。"
            "遥测每5秒上报temperature_c、humidity_pct、illuminance_lux、co2_ppm及执行器状态；"
            "无效传感器值使用null和invalid质量标记，不以模拟数值替代真实采集。"
        ),
        118: (
            "普通ESP32工程采用PlatformIO和Arduino框架。src/main.cpp负责安全启动与主循环调度；"
            "wifi和mqtt模块负责连接、订阅、发布和确认；Dht11Sensor、BH1750、JW01_CO2、PanTilt与"
            "WaterGunController分别封装传感器和执行器。platformio.ini固定esp32dev开发板、Adafruit"
            "DHT、FastLED与ArduinoJson等依赖，使构建环境可复现，JSON处理由ArduinoJson完成[6]。"
        ),
        120: (
            "普通ESP32固件已通过PlatformIO构建，后端pytest共165项测试及7项子测试通过，Web生产"
            "构建通过。软件链路已经覆盖真实EMQX通信、遥测持久化、REST/SSE、命令队列、ACK处理、"
            "静态与动态水枪、历史查询、摄像头、知识库和大模型服务。实物交付仍需按最终结构校准"
            "喷嘴方向、垂直落点、水泵功率和摄像头参数；这些标定影响命中精度，但不改变统一通信和"
            "安全控制流程。"
        ),
        124: (
            "舵机使用独立稳定5 V电源并与ESP32共地，水泵和补光灯经匹配驱动模块连接，GPIO不得直接"
            "带载。软件通过命令有效期、输入范围、先停泵后转向、断网急停、定时截止、动态超时、"
            "同目标去重和最终ACK降低误动作风险。Wi-Fi密码、EMQX账号、CA证书和模型密钥均保存在"
            "Git忽略的本机配置中；Web不接触设备级凭据。"
        ),
        127: (
            "系统分别定义设备能力、遥测、命令、状态和确认，新增硬件时先声明能力，再固定字段、"
            "单位、范围、错误码和验收方式。QoS 1允许重复投递，因此普通ESP32按command_id幂等，"
            "后端按设备确认更新命令状态。传感器无效时保留字段并显式标记质量，不以0冒充有效数据；"
            "旧farm/test主题、模拟字段和本地Mosquitto均不再参与正式运行。"
        ),
        130: "完成不同植物与季节条件下的温湿度、光照、二氧化碳目标范围配置，持续丰富家庭植物知识库。",
        131: "结合最终喷嘴高度、水压和安装方向标定距离—俯仰角—泵功率关系，提高定点浇灌落点一致性。",
        132: "完善C5语音与屏幕在弱网、断网和模型服务异常时的本地提示，并扩展多轮养护问答体验。",
        133: "在保持用户确认和执行器安全边界的前提下，增加可解释的自动补光、定时养护和异常提醒策略。",
        137: (
            "[3] ASAIR. DHT11 Temperature and Humidity Sensor Module[EB/OL]. "
            "[2026-07-25]. https://www.aosong.com/en/Products/info.aspx?itemid=2257."
        ),
    }

    for index, text in replacements.items():
        replace_paragraph_text(document.paragraphs[index], text)

    # Static TOC entries and body headings use paired hyperlink/bookmark nodes.
    # Update only their visible runs so navigation remains intact.
    replace_toc_hyperlink_text(document.paragraphs[38], "3.4.1 语音与屏幕协同方式")
    replace_bookmarked_heading_text(document.paragraphs[98], "3.4.1 语音与屏幕协同方式")
    replace_toc_hyperlink_text(document.paragraphs[58], "5.3.1 应用优化方向")
    replace_bookmarked_heading_text(document.paragraphs[129], "5.3.1 应用优化方向")

    for paragraph_index in (67, 88, 93, 108, 113, 118):
        add_internal_citation_links(document.paragraphs[paragraph_index])
    ensure_paragraph_bookmark(document, document.paragraphs[137], "ref_3")

    # Keep every body heading with the paragraph that follows it. The original
    # template allowed level-three headings to remain alone at a page bottom.
    for paragraph in document.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
    # These two long sections otherwise leave their level-three heading as the
    # final line of the preceding page in the supplied competition template.
    document.paragraphs[63].paragraph_format.page_break_before = True
    document.paragraphs[89].paragraph_format.page_break_before = True
    document.paragraphs[90].paragraph_format.page_break_before = False

    document.save(input_path)
    patch_bibliography(input_path)


if __name__ == "__main__":
    main()
