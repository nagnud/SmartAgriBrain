from __future__ import annotations

import re
import shutil
import sys
import uuid
from copy import deepcopy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from refined_figures import generate_all as generate_refined_figures


BIB_NS = "http://schemas.openxmlformats.org/officeDocument/2006/bibliography"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def replace_paragraph_text(paragraph, text: str) -> None:
    first_rpr = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if first_rpr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, first_rpr)


def remove_after(document: Document, paragraph_index: int) -> None:
    body = document._element.body
    keep = set(p._p for p in document.paragraphs[: paragraph_index + 1])
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        if child not in keep:
            body.remove(child)


def next_bookmark_id(document: Document) -> str:
    ids = []
    for element in document._element.body.xpath(".//w:bookmarkStart"):
        value = element.get(qn("w:id"))
        if value and value.isdigit():
            ids.append(int(value))
    return str(max(ids, default=-1) + 1)


def bookmark_paragraph(document: Document, paragraph, name: str) -> None:
    bookmark_id = next_bookmark_id(document)
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bookmark_id)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bookmark_id)
    insert_at = 1 if paragraph._p.find(qn("w:pPr")) is not None else 0
    paragraph._p.insert(insert_at, start)
    paragraph._p.append(end)


def add_hyperlink(paragraph, text: str, anchor: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rstyle = OxmlElement("w:rStyle")
    rstyle.set(qn("w:val"), "Hyperlink")
    rpr.append(rstyle)
    run.append(rpr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_citation_links(paragraph) -> None:
    text = paragraph.text
    matches = list(re.finditer(r"\[([1-9][0-9]*)\]", text))
    if not matches:
        return
    first_rpr = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)

    def append_text(value: str, parent, linked: bool = False) -> None:
        if not value:
            return
        run = OxmlElement("w:r")
        if first_rpr is not None:
            run.append(deepcopy(first_rpr))
        if linked:
            rpr = run.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                run.insert(0, rpr)
            rstyle = OxmlElement("w:rStyle")
            rstyle.set(qn("w:val"), "Hyperlink")
            rpr.insert(0, rstyle)
        node = OxmlElement("w:t")
        node.text = value
        run.append(node)
        parent.append(run)

    cursor = 0
    for match in matches:
        append_text(text[cursor : match.start()], paragraph._p)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("w:anchor"), f"ref_{match.group(1)}")
        hyperlink.set(qn("w:history"), "1")
        append_text(match.group(0), hyperlink, True)
        paragraph._p.append(hyperlink)
        cursor = match.end()
    append_text(text[cursor:], paragraph._p)


def set_run_font(run, east_asia: str = "宋体", ascii_font: str = "Times New Roman", size: float = 10.5) -> None:
    run.font.name = ascii_font
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def format_body(paragraph, first_line: bool = True) -> None:
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.5
    if first_line:
        paragraph.paragraph_format.first_line_indent = Pt(21)
    for run in paragraph.runs:
        set_run_font(run)


def shade_paragraph(paragraph, fill: str = "F2F4F5") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def add_body(document: Document, text: str):
    paragraph = document.add_paragraph(text)
    format_body(paragraph)
    add_citation_links(paragraph)
    return paragraph


def add_bullet(document: Document, text: str):
    try:
        paragraph = document.add_paragraph(text, style="List Bullet")
    except KeyError:
        paragraph = document.add_paragraph("• " + text)
    paragraph.paragraph_format.left_indent = Pt(21)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    format_body(paragraph, first_line=False)
    add_citation_links(paragraph)
    return paragraph


def add_code(document: Document, text: str):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(12)
    paragraph.paragraph_format.right_indent = Pt(12)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.05
    shade_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, east_asia="微软雅黑", ascii_font="Consolas", size=8.5)
    return paragraph


def font(size: int, bold: bool = False):
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size=size, index=0)
    return ImageFont.load_default()


def rounded_box(draw, xy, title, lines, fill, outline="#23343B"):
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = xy
    draw.text(((x1 + x2) / 2, y1 + 24), title, font=font(31, True), fill="#102027", anchor="ma")
    y = y1 + 72
    for line in lines:
        draw.text(((x1 + x2) / 2, y), line, font=font(22), fill="#26383E", anchor="ma")
        y += 37


def arrow(draw, start, end, label=""):
    draw.line([start, end], fill="#006B5D", width=7)
    x2, y2 = end
    x1, y1 = start
    angle_x = -1 if x2 > x1 else 1
    angle_y = -1 if y2 > y1 else 1
    if abs(x2 - x1) >= abs(y2 - y1):
        draw.polygon([(x2, y2), (x2 + 22 * angle_x, y2 - 13), (x2 + 22 * angle_x, y2 + 13)], fill="#006B5D")
    else:
        draw.polygon([(x2, y2), (x2 - 13, y2 + 22 * angle_y), (x2 + 13, y2 + 22 * angle_y)], fill="#006B5D")
    if label:
        draw.text(((x1 + x2) / 2, (y1 + y2) / 2 - 18), label, font=font(20), fill="#004D40", anchor="ms")


def create_architecture(path: Path) -> None:
    image = Image.new("RGB", (1800, 1050), "#F7F9F8")
    draw = ImageDraw.Draw(image)
    draw.text((900, 35), "SmartAgriBrain端到端系统架构", font=font(42, True), fill="#163238", anchor="ma")
    rounded_box(draw, (65, 155, 445, 430), "家庭现场设备", ["DHT11 / BH1750 / JW01", "ESP32采集与执行", "补光灯 / 二维水枪"], "#E2F3EA")
    rounded_box(draw, (65, 610, 445, 890), "交互终端", ["ESP32-C5语音与显示", "PC外接摄像头", "家庭用户"], "#FFF2D8")
    rounded_box(draw, (650, 335, 1110, 700), "云端与服务层", ["EMQX MQTT TLS", "FastAPI领域服务", "SQLite数据与知识库", "大模型 / 语音 / 视觉"], "#DDECF7")
    rounded_box(draw, (1350, 155, 1735, 430), "Web管理端", ["实时状态 / 历史曲线", "报警 / 设备控制", "知识库 / AI建议"], "#F4E5EE")
    rounded_box(draw, (1350, 610, 1735, 890), "外部服务", ["天气与农业公开来源", "云端大模型", "语音识别与合成"], "#ECE8F7")
    arrow(draw, (445, 295), (650, 420), "MQTT")
    arrow(draw, (650, 505), (445, 360), "命令/ACK")
    arrow(draw, (445, 750), (650, 620), "HTTPS/MQTT")
    arrow(draw, (1110, 420), (1350, 295), "REST/SSE")
    arrow(draw, (1350, 360), (1110, 505), "用户意图")
    arrow(draw, (1110, 620), (1350, 750), "HTTPS")
    draw.text((900, 995), "浏览器不直连Broker；后端是设备命令的唯一发布者，设备ACK才代表输出已应用。", font=font(22), fill="#38494E", anchor="ma")
    image.save(path, optimize=True)


def create_message_flow(path: Path) -> None:
    image = Image.new("RGB", (1800, 1050), "#FAFAF7")
    draw = ImageDraw.Draw(image)
    draw.text((900, 35), "遥测与控制的双向消息流", font=font(42, True), fill="#163238", anchor="ma")
    columns = [(70, "ESP32"), (490, "EMQX"), (900, "FastAPI/SQLite"), (1370, "Web/C5")]
    for x, name in columns:
        draw.rounded_rectangle((x, 120, x + 300, 210), radius=12, fill="#E4EFEA", outline="#315A50", width=3)
        draw.text((x + 150, 148), name, font=font(28, True), fill="#17372F", anchor="ma")
        draw.line((x + 150, 210, x + 150, 965), fill="#9AABA5", width=3)
    steps = [
        (300, 220, 640, "telemetry：每5秒，QoS 1"),
        (430, 640, 1050, "校验、去重、持久化"),
        (555, 1050, 1520, "SSE telemetry / REST state"),
        (690, 1520, 1050, "用户确认后的REST控制"),
        (800, 1050, 640, "command：QoS 1，不保留"),
        (900, 640, 220, "设备校验并执行"),
    ]
    for y, x1, x2, label in steps:
        arrow(draw, (x1, y), (x2, y), label)
    draw.text((900, 1000), "command_ack沿相反方向返回；command_id贯穿REST记录、MQTT命令、设备ACK和页面状态。", font=font(22), fill="#38494E", anchor="ma")
    image.save(path, optimize=True)


def create_watergun_flow(path: Path) -> None:
    image = Image.new("RGB", (1800, 1080), "#F8FAFB")
    draw = ImageDraw.Draw(image)
    draw.text((900, 35), "二维水枪安全执行状态机", font=font(42, True), fill="#163238", anchor="ma")
    items = [
        ("1 接收命令", "精确Topic、完整JSON、QoS 1队列"),
        ("2 协议校验", "command_id、expires_at、范围、会话序号"),
        ("3 先停水泵", "GPIO26立即写0，清除旧定时保护"),
        ("4 计算目标", "方位-90..90映射水平0..180；距离映射俯仰60..5"),
        ("5 移动舵机", "GPIO27水平、GPIO13俯仰，限位后写50 Hz PWM"),
        ("6 稳定等待", "非阻塞等待800 ms，主循环继续维护网络"),
        ("7 建立保护", "定时截止或动态3秒设备可见保活"),
        ("8 开泵并ACK", "GPIO26写实际PWM；executed不等于水流反馈"),
    ]
    y = 130
    for index, (title, detail) in enumerate(items):
        fill = "#E4F2EC" if index < 3 else "#E7EFF8" if index < 6 else "#FFF0D7"
        draw.rounded_rectangle((250, y, 1550, y + 88), radius=14, fill=fill, outline="#35545A", width=3)
        draw.text((300, y + 25), title, font=font(27, True), fill="#17363C", anchor="la")
        draw.text((650, y + 25), detail, font=font(22), fill="#2B4045", anchor="la")
        if index < len(items) - 1:
            arrow(draw, (900, y + 88), (900, y + 118))
        y += 118
    image.save(path, optimize=True)


CONTENT = [
    ("h1", "1 设计需求分析"),
    ("h2", "1.1 产品定位与应用场景"),
    ("h3", "1.1.1 家庭智能植物管家的定位"),
    ("p", "SmartAgriBrain定位为面向家庭阳台、室内种植空间及小型庭院的智能植物养护终端。产品不是对工业温室控制柜的简单缩小，而是围绕家庭用户“看得懂、摆得下、容易用、动作可控”的需求，采用方形一体化“小型智慧种植屋”结构，把环境感知、补光、二维定点灌溉、屏幕显示、Web远程管理和语音问答集中到同一设备中。"),
    ("p", "系统服务对象既包括缺少专业经验、希望降低日常养护负担的普通家庭，也包括需要观察植物环境变化、验证浇灌策略的学生和创客。设备持续形成环境状态，用户可以在现场通过显示屏快速查看，也可以在异地通过Web了解变化；当出现叶片异常、温湿度不适宜或养护知识不足时，还可以通过ESP32-C5语音终端调用云端大模型和植物知识库获得解释与建议。"),
    ("image", "product"),
    ("caption", "图1 SmartAgriBrain小型智慧种植屋产品概念"),
    ("h3", "1.1.2 典型使用场景"),
    ("p", "阳台场景中，光照受朝向、季节和天气影响明显，用户往往只能在早晚观察植物。本系统通过光照、空气温湿度和二氧化碳浓度的连续采样补足人工观察间隙，并用历史曲线呈现一天内的变化。补光灯采用百分比PWM控制，便于在阴天、傍晚或室内环境中调整亮度，而不是只有简单开关。"),
    ("p", "室内种植场景中，植物可能摆放在书房、客厅或封闭种植架内。设备需要低占地、低操作门槛，并避免持续喷水、舵机越界和网络失联等风险。二维水枪以目标方位和地面距离描述喷水位置，控制器内部统一完成停泵、转向、等待和开泵，减少用户直接操作多个硬件参数的负担。"),
    ("p", "小型庭院场景中，用户更关注远程查看、定时操作、摄像头观察和异常提醒。当前原型使用连接后端主机的PC外接摄像头提供MJPEG预览与抓帧分析；摄像头是视觉观察与定位的输入，不直接代替传感器，也不绕过用户确认触发高风险喷水动作。"),
    ("h2", "1.2 用户问题与需求边界"),
    ("h3", "1.2.1 需要解决的主要问题"),
    ("p", "第一，家庭养护数据分散且不连续。人工只能获得某一时刻的主观感受，难以判断温度、湿度和光照是短时波动还是长期偏离。系统需要以统一单位采集、上报、保存和展示数据，并区分正常、过期、无效和未支持状态，避免把裸板默认值误认为真实环境。"),
    ("p", "第二，传统浇水通常只能控制总量，不能准确描述希望浇到的区域。系统需要把Web上选择的位置转换为水平轴和垂直轴目标，同时保证水泵与舵机动作顺序安全。位置没有变化时不应重复写舵机PWM；动态模式也不能因为页面心跳或消息重发而在目标角与初始角之间往复。"),
    ("p", "第三，设备信息、控制入口和养护知识彼此割裂。系统需要让同一站点的实时数据、历史记录、摄像头画面、报警、知识库和助手对话在统一Web中协同，并允许C5语音终端复用后端上下文。控制建议与真实执行必须分离，大模型只能给出建议或待确认动作，不能直接写GPIO。"),
    ("h3", "1.2.2 当前产品边界"),
    ("p", "当前现场执行器只有补光灯和二维云台水枪，不包含风机、加热器、制冷器、卷帘或独立水泵设备。GPIO26控制的水泵负责把水输送到水枪，属于水枪内部部件；GPIO27和GPIO13的两个SG90同样属于水枪内部部件。正式协议只允许补光灯使用原子set命令，喷水必须通过target_position复合命令。"),
    ("p", "当前环境采集使用GPIO4上的DHT11获得空气温度与空气相对湿度，使用BH1750获得光照强度，使用JW01串口模块获得二氧化碳浓度。旧土壤湿度传感器和DS18B20不属于当前主固件采集路径，因此正文、遥测字段和验收项不再把土壤湿度列为已实现功能。"),
    ("p", "ESP32-C5承担语音和显示终端职责，后端已经提供C5设备标识、助手回复主题、语音识别和语音合成配置，但当前工作区没有纳入C5终端固件源码。本文可以描述已确定的交互职责和后端通道，不能把缺少源码和真机记录的C5固件写成已经在本仓库构建验收。"),
    ("h2", "1.3 功能需求"),
    ("h3", "1.3.1 环境监测与可视化"),
    ("p", "普通ESP32应在不阻塞网络和执行器状态机的前提下周期推进传感器驱动。DHT11物理采样间隔不低于2秒，遥测发布周期为5秒；同一次DHT11读取同时形成temperature_c和humidity_pct，避免为两个字段重复访问传感器。BH1750使用I2C地址0x23，JW01通过其串口驱动更新。"),
    ("p", "后端接收遥测后应验证Topic设备ID、message_id、采样时间、字段结构、单位和quality，再进行QoS 1去重、站点归属映射和SQLite持久化。Web需要显示最新值、设备在线状态、数据质量、历史趋势和报警记录，并用SSE获得实时事件，用REST完成断线后的状态校准。"),
    ("h3", "1.3.2 补光与精准灌溉"),
    ("p", "补光灯通过GPIO14高电平有效PWM驱动，频率1 kHz、分辨率8位，协议值域为0至90%。Web或语音助手提交用户意图后，由后端生成唯一command_id和约30秒有效期的MQTT命令；ESP32验证后写入实际占空比并返回ACK。"),
    ("p", "水枪应支持静态、动态和定时喷水。静态模式对一个位置执行一次目标设置，可持续喷水或按本地截止时间停止；动态模式使用session_id和严格递增sequence更新位置，并在设备3秒内没有收到合法新序号时停泵；定时模式同时使用UTC截止时间与ESP32单调时钟，后端停止消息只是第二重保险。"),
    ("h3", "1.3.3 Web、摄像头与智能助手"),
    ("p", "Web需要覆盖首页总览、实时监测、历史曲线、设备控制、病害图片、AI农事建议、知识库管理和报警记录八类业务视图。页面不能持有MQTT凭据，也不能把本地状态变化当成设备执行成功；命令从提交、排队、发布到ACK的生命周期应由后端记录并向页面发布。"),
    ("p", "后端摄像头服务负责独占打开PC外接摄像头，提供状态、配置、单帧JPEG和MJPEG流。视觉分析结果可以为目标位置、病害识别和养护建议提供上下文，但识别失败、摄像头断开或标定缺失时必须显式返回不可用状态，不得使用模拟画面或伪造结果。"),
    ("p", "智能助手应组合最新环境、历史趋势、天气、摄像头分析和知识库检索结果形成回答。知识型问题可以直接回复；补光和喷水等实体动作先生成待确认action，只有用户确认后才进入站点命令服务。助手输出应保留依据与风险提示，不能把语言模型生成内容视为传感器事实。"),
    ("h2", "1.4 非功能需求"),
    ("h3", "1.4.1 可靠性与一致性"),
    ("p", "家庭Wi-Fi可能短时断开，MQTT QoS 1也允许重复投递。设备因此采用持久会话、命令有效期、完整command_id幂等、内容指纹和ACK缓存；后端以message_id去重遥测，以command_id关联命令与ACK。QoS 1保证至少一次送达，不保证业务只执行一次，幂等处理是物理控制安全的一部分[1]。"),
    ("p", "浏览器采用SSE接收实时事件，同时每5秒通过REST校准站点状态，每60秒拉取历史数据。SSE断线时页面保留最后状态，但必须标记连接异常；重连后重新读取站点快照和共享助手会话，避免把断线期间遗漏的事件当成当前事实。"),
    ("h3", "1.4.2 安全、隐私与可维护性"),
    ("p", "水泵、灯和舵机不能由ESP32 GPIO直接带载。舵机使用独立稳定5 V电源并与ESP32共地，水泵和补光灯经过匹配的MOS管、继电器或驱动板。软件在启动、断网、命令非法、动态超时和定时结束时优先停泵，机械角度始终经过限位。"),
    ("p", "Wi-Fi密码、EMQX账号、CA路径、云端模型密钥和语音服务令牌只保存在本机忽略配置中，不写入前端源码、README或Git历史。浏览器只访问FastAPI；EMQX ACL应按后端、普通ESP32和C5分别授权最小Topic范围，防止某一端泄露后读取或发布无关设备主题。"),
    ("h3", "1.4.3 设计目标与验收原则"),
    ("bullet", "形成从真实传感器到EMQX、后端数据库、Web/C5展示的完整数据闭环，并明确每个字段的单位、质量和更新时间。"),
    ("bullet", "形成从Web或语音意图到后端命令、MQTT下发、ESP32执行和command_ack返回的完整控制闭环。"),
    ("bullet", "保证水枪目标变化时先停泵再转向，目标不变时不重复驱动舵机，断网、超时和定时结束均可独立停泵。"),
    ("bullet", "将源码存在、编译通过、自动化测试通过、真机动作通过和端到端验收通过作为不同结论记录，避免以软件日志替代物理结果。"),

    ("h1", "2 特色与创新"),
    ("h2", "2.1 小型智慧种植屋一体化设计"),
    ("h3", "2.1.1 从单点模块到家庭终端"),
    ("p", "本项目将传感器、补光、定点灌溉、屏幕、语音和云端服务组织为一个面向家庭的完整产品，而不是多个开发板示例的简单拼接。方形一体化结构使环境感知区域、植物生长区域、补光区域和喷水机构具有明确空间关系，用户能够把显示信息与眼前植物对应起来。"),
    ("p", "一体化并不意味着所有任务由同一芯片承担。系统按照实时性、外设资源和交互复杂度拆分职责：普通ESP32贴近传感器和执行器，保证安全状态机持续运行；ESP32-C5贴近用户，承担显示、录音和播放；后端负责计算、存储、知识和跨端协调。"),
    ("h2", "2.2 双MCU与云边协同"),
    ("h3", "2.2.1 普通ESP32与C5职责分离"),
    ("p", "普通ESP32只处理需要确定时序和硬件约束的工作，包括DHT11、BH1750、JW01采集，GPIO14补光PWM，GPIO26水泵PWM以及GPIO27/GPIO13双SG90控制。即使云端不可用，设备仍能执行本地定时截止、动态超时和网络断开停泵。"),
    ("p", "C5终端面向语音和显示，接收后端整理后的站点状态与助手回复。它不直接订阅普通ESP32的GPIO命令，也不保存执行器控制规则。这样可以让屏幕界面、语音识别或大模型调用发生故障时，不破坏现场ESP32的安全动作。"),
    ("image", "modules"),
    ("caption", "图2 现场设备、双MCU、摄像头与云端模块协同关系"),
    ("h3", "2.2.2 云端负责知识与审计"),
    ("p", "后端把EMQX消息转换为站点领域对象，保存遥测、设备在线状态、能力、命令、ACK、报警、知识库和助手会话。Web和C5看到的是同一站点状态，不需要各自解释原始MQTT报文；控制来源、创建时间、拒绝原因和最终ACK也可以在服务端统一审计。"),
    ("h2", "2.3 面向位置的二维精准水枪"),
    ("h3", "2.3.1 用目标位置代替独立舵机命令"),
    ("p", "用户关心的是“向哪个位置喷水”，而不是GPIO27写多少脉宽。系统因此使用bearing_deg和ground_range_mm表达目标，以正前方为0度、左侧为负、右侧为正；后端与设备共同验证范围，设备再转换为水平和垂直舵机角度。水泵、水平轴和垂直轴不作为三个可乱序发布的通用执行器。"),
    ("p", "水平映射把Web的-90至90度线性映射到舵机命令0至180度，0度方位对应水平舵机90度正前方。垂直机构的真实可用角度为5至60度：5度平行向前，60度垂直向上；当前控制模型把300至1200毫米地面距离线性映射为60至5度。该映射建立了可执行关系，但仍需结合喷嘴、水压和安装高度进行落点标定。"),
    ("h3", "2.3.2 状态机保证动作顺序"),
    ("p", "目标改变时，WaterGunController先把GPIO26写为0，再计算并写入两个舵机目标，非阻塞等待约800毫秒后才根据喷射设置开泵。状态机运行期间主循环继续维护Wi-Fi、MQTT和传感器，不使用阻塞delay占用网络处理。"),
    ("p", "相同位置的动态保活只更新会话序号和安全计时，不重复写舵机PWM；在第一次移动仍处于稳定等待时，新保活接管最终ACK但保留原截止时间。这一处理直接避免页面目标未变化时舵机因重复命令不断转动。"),
    ("h2", "2.4 统一通信契约"),
    ("h3", "2.4.1 Topic可推导信息与精简JSON"),
    ("p", "设备身份和协议版本由smartagribrain/v1/devices/{device_id}/...这一Topic层级确定，站点由后端设备注册关系确定，因此遥测JSON不再重复schema_version、device_id和site_id。精简后的消息保留去重、采样、传感器、执行器与连接质量等真正变化的数据，降低5秒周期上报的冗余。"),
    ("p", "五类业务Topic各自承担单一职责：telemetry是周期快照，status是保留的在线状态，capabilities是保留的能力声明，command是后端按需下发的控制，command_ack是设备对单条命令的最终确认。它们不是都要周期发送，只有telemetry固定每5秒发布。"),
    ("h3", "2.4.2 REST、SSE与MQTT使用同一领域数据"),
    ("p", "前端使用temperature_c、humidity_pct、illuminance_lux、co2_ppm等统一字段，后端数据库、SSE事件和MQTT解析都沿用相同含义。控制链路用command_id贯穿REST资源、数据库记录、MQTT命令和设备ACK，减少跨端联调时的字段翻译和状态歧义。"),
    ("h2", "2.5 知识库、视觉与大模型协同"),
    ("h3", "2.5.1 让建议具有现场上下文"),
    ("p", "助手不是只向通用模型发送一句问题，而是可以组合最新环境、历史趋势、天气、植物知识条目、病害图片和摄像头定位结果。后端知识库将内容划分为知识库、条目和片段，支持检索、维护和分析；在线农业来源也可以单独启停和测试。"),
    ("p", "涉及补光或喷水时，模型只生成候选动作及理由，用户确认后由后端按照正式契约创建命令。这个“建议、确认、执行、回执”链路把生成式模型的不确定性隔离在硬件控制之前，也使Web和C5能够显示同一个待确认动作。"),

    ("h1", "3 功能设计"),
    ("h2", "3.1 环境采集功能"),
    ("h3", "3.1.1 DHT11温湿度采集"),
    ("p", "DHT11数据线连接GPIO4，驱动对象在setup阶段调用begin初始化，在主循环中通过update推进采样。传感器读数使用isnan检查温度和湿度有效性；一次物理读取同时更新两个缓存值，避免两个领域字段分别触发总线访问。DHT11的测量与接口特性以制造商资料为依据[2]。"),
    ("p", "驱动把最小物理采样间隔设为2000毫秒，遥测周期为5000毫秒，因此一条遥测使用最近一次有效采样。采集失败时串口输出READ_FAIL，后续遥测应通过quality表达无效或过期，而不是用前一次值假装当前采样成功。"),
    ("h3", "3.1.2 光照与二氧化碳采集"),
    ("p", "BH1750使用I2C地址0x23并在连续高分辨率模式下工作，上报illuminance_lux。其数字输出避免ESP32端进行光敏电阻ADC换算，适合直接进入统一单位的历史曲线和补光判断[3]。"),
    ("p", "JW01二氧化碳模块由独立串口驱动更新，上报co2_ppm。CO2数据用于环境观察、报警与养护建议，不对应当前执行器；系统不存在CO2阀门或通风控制，因此后端不能根据该指标生成不存在的硬件命令。"),
    ("h3", "3.1.3 遥测质量与发布节奏"),
    ("p", "每条telemetry带UUID v4格式message_id和sampled_at，并包含四类传感器对象、执行器快照及Wi-Fi RSSI。quality只允许ok、stale、invalid和unsupported。设备收到数据不等于传感器已接入，裸板上的0、-127或其他占位值只能验证链路，真机验收必须同时核对传感器接线和串口日志。"),
    ("p", "5秒周期是环境变化速度、网络负载和数据库写入量之间的折中。后端把原始遥测保存后，以约10秒时间桶生成站点历史采样，前端按时间窗口压缩为图表点。这样实时卡片保持灵敏，历史曲线又不会因QoS重发或高频重复值无限增长。"),
    ("h2", "3.2 补光功能"),
    ("h3", "3.2.1 PWM控制与命令范围"),
    ("p", "补光灯连接GPIO14，LEDC通道0，频率1 kHz，分辨率8位，高电平有效。协议值0表示关闭，1至90表示目标百分比；固件把百分比换算到0至255的占空比。LEDC提供可配置频率与分辨率的硬件PWM能力[5]。"),
    ("p", "补光灯是当前唯一通用原子执行器，命令使用operation=set、target=grow_light和整数value。后端在能力声明未确认supported=true、设备离线或值越界时应拒绝发布；ESP32也执行第二次值域校验，形成云端与端侧双重防线。"),
    ("h3", "3.2.2 手动、建议与自动化边界"),
    ("p", "Web可由用户直接设置亮度，语音助手也可以生成待确认的补光动作。页面侧的智能托管算法属于需求计算，不是ESP32正式命令；只有后端把最终需求转换为grow_light set命令后，才会进入真实MQTT链路。"),
    ("p", "关闭托管或发生异常时，系统应显式发布grow_light=0，而不是仅在页面把滑块归零。后端等待设备ACK后才能把状态标记为已应用；没有电流或照度闭环传感器时，ACK只证明GPIO占空比已经写入。"),
    ("h2", "3.3 二维水枪功能"),
    ("h3", "3.3.1 静态模式"),
    ("p", "静态模式适合用户选择一个固定位置后执行一次转向。请求包含ground_range_mm、bearing_deg、spray_enabled、pump_control_percent、spray_schedule和可选定时信息。continuous表示持续到下一条停止命令，timed表示由设备本地计时自动结束。"),
    ("p", "停止喷水仍然是一条合法target_position命令，spray_enabled=false时设备保持目标但立即把水泵设为0，并清除旧定时保护。Web需要将停止结果与最后目标分别显示，不能通过把位置恢复为默认值来表达停止。"),
    ("h3", "3.3.2 动态模式"),
    ("p", "动态模式先由后端创建session_id，再按递增sequence更新目标。目标位置更新与页面heartbeat是不同概念：未喷水时，页面heartbeat只维护后端会话，不应反复向设备发送相同位置；喷水时，后端需要生成设备可见的新序号保活，使ESP32的3秒超时可以判断云端链路仍然有效。"),
    ("p", "ESP32按同一session_id下的sequence拒绝旧目标。若消息因QoS 1重发，command_id和内容相同则重发缓存ACK；若同一command_id被用于不同内容，则按INVALID_COMMAND拒绝。目标未变化时控制器跳过舵机写入，只刷新动态保护或喷射状态。"),
    ("h3", "3.3.3 舵机角度与位置关系"),
    ("p", "水平轴GPIO27的Web方位与真实舵机命令关系为：-90度对应0度、-45度对应45度、0度对应90度、45度对应135度、90度对应180度。软件完整使用0至180度命令范围，但实体安装方向、连杆干涉和SG90脉宽端点必须在断泵状态下逐步标定[4]。"),
    ("p", "垂直轴GPIO13受机构限制，只允许5至60度。当前线性模型中300毫米对应60度，1200毫米对应5度，每增加100毫米角度约降低6.11度。这个关系描述的是控制模型而非弹道真值；水流落点还受喷嘴高度、泵压、软管阻力和重力影响。"),
    ("image", "coordinates"),
    ("caption", "图3 Web目标位置与水平、垂直舵机角度映射"),
    ("h3", "3.3.4 水泵功率与安全联锁"),
    ("p", "GPIO26水泵使用LEDC通道1、1 kHz、8位PWM，协议范围0至100%。当前临时规则把1至39%的非零请求提升到40%，用于避免假定的低占空比无法启动；该阈值尚未经过真实水泵测量，ACK必须返回actual_value而不是原请求值。"),
    ("p", "任何目标变化、网络断开、非法水枪命令、定时结束或动态超时都进入统一停泵路径。舵机移动前停泵可以降低扫过非目标区域时误喷的风险；舵机供电异常、泵驱动故障或软管漏水仍需硬件保险与人工急停，软件状态机不能替代电气保护。"),
    ("h2", "3.4 MQTT通信功能"),
    ("h3", "3.4.1 五类设备Topic"),
    ("bullet", "command：smartagribrain/v1/devices/greenhouse_001_s3/command。后端发布、ESP32订阅，QoS 1、retain=false，只在用户操作、自动策略或助手确认产生真实命令时触发。"),
    ("bullet", "command_ack：smartagribrain/v1/devices/greenhouse_001_s3/command_ack。ESP32发布、后端订阅，QoS 1、retain=false，每条命令最终产生一次结果，重复命令可重发缓存结果。"),
    ("bullet", "telemetry：smartagribrain/v1/devices/greenhouse_001_s3/telemetry。ESP32每5秒发布，QoS 1、retain=false，承载传感器、执行器和RSSI快照。"),
    ("bullet", "status：smartagribrain/v1/devices/greenhouse_001_s3/status。连接或重连时发布online，异常断开由LWT发布offline，QoS 1、retain=true。"),
    ("bullet", "capabilities：smartagribrain/v1/devices/greenhouse_001_s3/capabilities。连接或重连时发布传感器、补光灯和水枪能力，QoS 1、retain=true。"),
    ("image", "mqtt"),
    ("caption", "图4 EMQX五类设备Topic、方向与触发机制"),
    ("h3", "3.4.2 EMQX连接和权限"),
    ("p", "项目只使用外部EMQX，不使用本地Mosquitto。普通ESP32和后端通过mqtts://主机:8883连接，启用TLS证书校验与持久会话。后端账号只需要订阅设备上行并发布command；普通ESP32账号只需要订阅自身command并发布自身四类上行Topic。Broker部署、认证和Topic权限依据EMQX官方文档配置[10]。"),
    ("p", "status和capabilities使用retain，使后端重连后立即获得最后快照；command绝不保留，避免新设备上线后执行历史动作。MQTT 3.1.1没有消息过期属性，因此expires_at由ESP32强制验证，已过期或异常超前超过300秒的命令不会写GPIO[1]。"),
    ("h3", "3.4.3 紧凑遥测示例"),
    ("code", """{
  "message_id": "8d42b9ca-9e98-46f4-8f73-64f9c223e979",
  "sampled_at": 1784966405000,
  "sensors": {
    "temperature_c": {"value": 26.3, "quality": "ok"},
    "humidity_pct": {"value": 61.0, "quality": "ok"},
    "illuminance_lux": {"value": 842.0, "quality": "ok"},
    "co2_ppm": {"value": 527.0, "quality": "ok"}
  },
  "actuators": {"grow_light_pct": 40},
  "water_gun": {"pump_pct": 0, "pan_deg": 90, "tilt_deg": 5,
                "active": false, "timed": false, "dynamic": false},
  "connection": {"rssi_dbm": -56}
}"""),
    ("p", "该JSON不重复Topic已经提供的协议版本和设备ID，也不重复后端设备注册能够确定的site_id。message_id用于QoS 1去重，sampled_at用于判断数据时效；传感器对象始终携带quality，执行器值来自统一硬件状态快照而不是从最近一次命令反推。"),
    ("h3", "3.4.4 水枪命令示例"),
    ("code", """{
  "command_id": "c734a4f2-73ef-4cc0-9426-35e63477a7f1",
  "expires_at": 1784966430000,
  "command": {
    "operation": "target_position",
    "target": "position",
    "value": 1,
    "water_gun": {
      "ground_range_mm": 650,
      "bearing_deg": -25,
      "mode": "static",
      "spray_enabled": true,
      "pump_control_percent": 55,
      "session_id": null,
      "sequence": 0,
      "spray_schedule": "timed",
      "spray_duration_seconds": 12,
      "spray_ends_at": 1784966417000
    }
  }
}"""),
    ("p", "命令只包含设备执行需要的字段。source、site_id、issued_at和用户原因由后端数据库审计，不在MQTT报文中重复。动态模式必须提供非空session_id和大于0的递增sequence；静态模式不携带动态会话。定时模式的duration和ends_at必须同时有效。"),
    ("h2", "3.5 后端服务功能"),
    ("h3", "3.5.1 设备接入与状态聚合"),
    ("p", "FastAPI启动时创建数据库表、监控运行时和MQTT运行时。MQTT服务订阅telemetry、status、capabilities和command_ack，按Topic提取device_id，校验JSON后调用站点服务。站点服务把普通ESP32映射到greenhouse_001，生成统一SiteState并通过事件总线发布SSE。"),
    ("p", "后端命令服务在数据库建立pending记录，再由MQTT传输发布。ACK到达后按command_id更新为executed或rejected，并保存actual、error和feedback_verified。超时没有收到ACK时标记timed_out，页面不能继续显示“执行成功”。"),
    ("h3", "3.5.2 REST与SSE接口"),
    ("p", "站点核心接口包括GET /api/v1/sites/{site_id}/state、GET /history、GET /events和POST /commands。水枪提供状态、preview、static、dynamic/start、dynamic/target、dynamic/heartbeat、dynamic/spray和dynamic/stop等接口。API按业务意图组织，不直接暴露GPIO编号。"),
    ("p", "SSE事件包括telemetry、device_status、capabilities、command_update、water_gun、assistant_message和assistant_action。页面丢失SSE后通过REST读取完整快照；SSE用于降低刷新延迟，不承担唯一数据存储职责。FastAPI与浏览器原生事件流适合单向实时状态推送[7]。"),
    ("h3", "3.5.3 SQLite数据持久化"),
    ("p", "默认数据库是backend_api/smartagribrain.db，使用SQLAlchemy管理，无需独立启动数据库服务。设备数据包括telemetry_records、telemetry_receipts、device_presence、device_commands、site_snapshots、site_history_samples和site_commands；业务数据还包括alarms、app_states、knowledge_bases、knowledge_items、knowledge_chunks、disease_photos以及助手会话和动作。"),
    ("p", "SQLite以单文件形式降低家庭原型部署成本，后端在同一事务中更新命令或站点状态，适合单机演示和小规模数据量[9]。若后续多实例部署或遥测规模增长，可通过DATABASE_URL切换PostgreSQL，但领域模型、REST和MQTT契约不应随数据库更换而改变。"),
    ("h3", "3.5.4 摄像头、视觉与天气"),
    ("p", "摄像头服务统一拥有PC外接摄像头句柄，提供/api/v1/camera/status、/config、/frame.jpg和/stream.mjpg。集中持有可以避免前端、分析任务和多个进程同时打开同一设备。关闭后端时必须释放OpenCV VideoCapture，否则即使Web已关闭，系统仍会显示摄像头被占用[11]。"),
    ("p", "视觉服务支持病害图片上传、保存、分析和删除，也支持摄像头抓帧后的作物位置分析。天气服务提供当前天气与组合预报，农业来源服务管理在线来源的启停与连通性。所有外部服务失败都应局部降级，不能阻断设备遥测和安全控制。"),
    ("h2", "3.6 Web管理端功能"),
    ("h3", "3.6.1 页面结构与状态同步"),
    ("p", "Web使用Vue 3、TypeScript、Vite和ECharts构建[8]。八个业务视图分别承担总览、实时监测、历史曲线、病害识别、AI建议、设备控制、知识库管理和报警记录。api.ts集中处理REST、SSE、上传、流式助手和摄像头资源URL，避免组件各自拼接接口。"),
    ("p", "页面状态同时写入localStorage和后端/api/v1/state/dashboard。浏览器本地保存用于后端暂时不可用时快速恢复，后端状态用于多会话同步。设备事实始终以后端站点快照为准，本地缓存只用于界面恢复。"),
    ("h3", "3.6.2 实时、历史与报警"),
    ("p", "首页和实时页接收SSE遥测后立即更新卡片，并把样本合并进历史缓存。历史页按5分钟至12小时窗口读取真实记录，可选择温度、湿度、光照和CO2指标。每60秒REST校准补回断线样本，相同时间槽保留较新记录。"),
    ("p", "报警服务根据持久化阈值和数据状态生成记录，Web显示开放、已确认等状态。用户确认报警只代表已阅读或已处理提示，不代表温度、湿度或设备故障已经恢复；报警恢复需要后续遥测和规则重新评估。"),
    ("h3", "3.6.3 水枪交互"),
    ("p", "控制页允许在可视区域指定方位和距离，预览接口只计算目标状态，不发布MQTT；用户确认后才调用static或dynamic接口。动态拖动进行节流，避免鼠标每个像素变化都生成命令；喷射启停与目标更新分开，便于先对准再喷水。"),
    ("p", "前端每500毫秒更新倒计时显示，但不向设备发送控制。页面heartbeat用于维持后端动态会话；是否向ESP32发送设备可见保活由后端根据正在喷水状态决定。离开动态模式、页面卸载或用户停止时，前端调用停止接口并清理定时器。"),
    ("h3", "3.6.4 助手和知识库"),
    ("p", "页面AI对话先POST /api/v1/assistant/turns创建任务，再通过SSE读取progress、completed或failed事件。共享助手会话使用站点接口保存Web与C5消息和待确认动作。涉及水枪和补光的action必须经过decision接口确认，随后才能创建设备命令。"),
    ("p", "知识库页面支持知识库和条目的新增、修改、删除与问答分析；后端把长文本拆成片段用于检索。病害图像、当前环境、天气和知识检索结果可以进入助手上下文，但前端需要区分传感器事实、外部资料、模型推断和控制建议。"),
    ("h2", "3.7 ESP32-C5语音显示功能"),
    ("h3", "3.7.1 显示职责"),
    ("p", "C5显示屏用于呈现当前环境参数、设备在线状态、补光与水枪状态、养护提醒和助手回复摘要。后端可通过smartagribrain/v1/devices/{C5设备ID}/view_state发布保留视图状态，使终端重连后迅速恢复；显示层不需要解析普通ESP32原始遥测。"),
    ("h3", "3.7.2 语音交互链路"),
    ("p", "语音链路由C5录音、后端接收音频、云端ASR转写、助手编排、知识与大模型生成、TTS合成以及C5播放组成。后端speech_service配置火山引擎语音识别和语音合成参数，assistant/response Topic用于推送非保留回复。"),
    ("p", "语音识别出的动作意图同样进入确认流程。补光、水枪等操作需要明确参数和用户确认；缺少位置、亮度或时长时，助手应提出补充问题并保留短时会话上下文。由于当前仓库没有C5固件源码，终端显示刷新、音频缓冲和实机时延需要在C5工程归档后单独验收。"),
    ("h2", "3.8 典型业务流程"),
    ("h3", "3.8.1 用户查看环境与历史"),
    ("p", "用户打开Web后，页面首先读取站点完整快照。后端从site_snapshots返回最新传感器、执行器、设备在线状态、能力和水枪会话，页面据此建立初始显示；随后建立SSE连接，接收新遥测和状态事件。若ESP32此时离线，页面仍可展示最后一次记录，但必须同时显示采样时间和离线状态，不能让历史值看起来像实时值。"),
    ("p", "当新的telemetry到达EMQX时，后端按Topic识别设备，验证UUID message_id和字段质量，忽略QoS 1产生的重复消息，更新站点快照并保存历史时间桶。SSE把同一标准化对象推给Web，卡片即时刷新；历史页下次校准时从数据库读取相同样本。这个流程保证实时视图与历史视图使用同一数据来源。"),
    ("p", "用户切换历史时间窗口时，Web只向后端请求相应小时范围，不从浏览器本地模拟曲线。后端按时间排序返回保存样本，页面根据画布宽度和指标选择压缩显示。温度单位固定为摄氏度、湿度为百分比、光照为勒克斯、二氧化碳为ppm，任何界面都不重复进行单位猜测。"),
    ("h3", "3.8.2 用户执行一次定点浇灌"),
    ("p", "用户在水枪控制区选择目标后，页面先调用preview接口得到后端认可的方位、距离和预计控制状态。预览不会创建MQTT命令，也不会让舵机动作。用户选择持续或定时策略、泵功率并确认风险提示后，页面才提交static请求。"),
    ("p", "后端检查站点、普通ESP32在线状态、能力、位置范围、喷射策略和定时参数，创建水枪状态与site_commands记录，再生成带command_id和expires_at的target_position报文。MQTT运行时以QoS 1、retain=false发布到设备command Topic，并把命令状态通过SSE更新为pending或published。"),
    ("p", "ESP32完成幂等和值域校验后先停泵，计算两个舵机角度并进入800毫秒非阻塞等待。稳定后才设置定时保护并写入泵PWM，随后发布executed ACK。后端保存actual值并更新水枪状态，Web最终显示“设备已应用控制输出”；如果ACK拒绝或超时，则显示错误且不把本地目标当作真实执行状态。"),
    ("h3", "3.8.3 用户进行动态跟踪"),
    ("p", "动态模式开始时，后端创建唯一session_id并初始化sequence。页面拖动目标时对高频指针事件进行节流，只在目标有意义变化时调用dynamic/target；后端为每次设备可见更新递增sequence。ESP32保存最后会话与序号，旧消息、重复序号和延迟到达目标都不能覆盖新位置。"),
    ("p", "未喷水时，页面heartbeat只用于后端判断页面仍在动态模式，不需要周期重发相同角度。开始喷水后，后端按安全周期发送带新sequence的同目标保活，设备只刷新lastDynamicCommandMs而不重复写舵机PWM。页面关闭、网络中断或后端超过会话期限时停止保活，ESP32最多3秒后独立停泵。"),
    ("p", "用户更新位置时，ESP32先关闭水泵再转动舵机，避免水束扫过路径；目标到位并等待稳定后才恢复喷水。用户点击停止时，页面调用dynamic/stop，后端发布spray_enabled=false并结束会话。即使停止命令丢失，设备侧动态超时仍提供最后保护。"),
    ("h3", "3.8.4 用户通过助手获取建议并执行"),
    ("p", "用户询问“现在是否需要补光”时，助手编排器读取最新光照、数据质量、采样时间、天气和知识库结果。若设备离线或光照数据过期，回答应先说明证据不足；若数据有效，则给出建议亮度及依据，但不立即发布命令。"),
    ("p", "当用户明确要求执行补光时，助手创建待确认动作，记录target、value、理由和失效时间。Web或C5展示动作摘要，用户确认后后端再次检查设备状态和参数，再走正常命令服务；取消、超时或二次检查失败均不会发布MQTT。"),
    ("p", "水枪语音请求还需要方位、距离、是否喷水、泵功率和时长。缺少参数时助手通过多轮会话追问，不使用未说明的默认位置。高风险动作确认后仍必须等待普通ESP32的最终ACK，语音回复中不能仅凭“命令已提交”播报“浇水完成”。"),

    ("h1", "4 系统实现"),
    ("h2", "4.1 总体架构与信息流"),
    ("h3", "4.1.1 系统分层"),
    ("p", "系统从下到上分为现场硬件层、嵌入式接入层、消息与服务层、数据与智能层以及用户交互层。现场硬件层包含传感器、灯、泵、舵机和摄像头；普通ESP32与C5形成嵌入式接入层；EMQX、FastAPI和SQLite承担消息、业务与持久化；Web、屏幕和语音构成交互入口。"),
    ("image", "architecture"),
    ("caption", "图5 SmartAgriBrain端到端系统架构"),
    ("h3", "4.1.2 遥测与控制闭环"),
    ("p", "遥测方向为传感器到ESP32、EMQX、FastAPI、SQLite，再由REST/SSE提供给Web和C5。控制方向从Web或C5的用户意图开始，经后端校验和记录后发布command，ESP32执行并返回command_ack，后端再把最终状态推送给交互端。"),
    ("image", "message"),
    ("caption", "图6 遥测、命令和回执的双向消息流"),
    ("h2", "4.2 普通ESP32硬件实现"),
    ("h3", "4.2.1 GPIO与LEDC资源"),
    ("p", "GPIO4连接DHT11数据线；BH1750占用I2C总线；JW01使用项目配置的串口。GPIO14绑定LEDC通道0，GPIO26绑定通道1，均使用1 kHz和8位分辨率。GPIO27水平舵机绑定通道2，GPIO13垂直舵机绑定通道3，均使用50 Hz和16位分辨率。通道分离防止后初始化模块覆盖既有PWM配置。"),
    ("p", "上电顺序固定为：启动串口，配置灯和泵PWM，将两者写为0，再把水平轴置于90度正前方、垂直轴置于5度平射位置，随后初始化传感器和网络。安全关闭状态先于Wi-Fi和MQTT，避免网络启动失败时执行器处于未知输出。"),
    ("image", "gpio"),
    ("caption", "图7 普通ESP32传感器、执行器与GPIO资源分配"),
    ("h3", "4.2.2 SG90脉宽与机械保护"),
    ("p", "PanTilt模块把角度限制在各自机械范围，再换算为50 Hz PWM脉宽。水平轴软件范围0至180度，垂直轴软件范围5至60度；初始化、目标映射、直接写入和遥测快照都使用同一限制常量，避免某一入口恢复旧的90度垂直初值。"),
    ("p", "SG90不是360度连续旋转舵机，不能通过超范围脉宽追求更大行程。真机标定应断开水泵，从中点开始逐步扩大水平端点，观察机械卡滞和电流；垂直轴只在已确认5至60度内测试。供电压降可能造成舵机仅转动约30度或反复复位，因此需要独立5 V供电和可靠共地。"),
    ("h2", "4.3 ESP32软件实现"),
    ("h3", "4.3.1 模块组织"),
    ("p", "src/main.cpp负责对象创建、初始化顺序和主循环调度；include/iot_contract.h保存设备ID、Topic前缀、固件版本、值域和标定常量；lib/mqtt负责TLS连接、Topic、命令解析、幂等、ACK和遥测；lib/WaterGun与lib/PanTilt封装水枪状态机和舵机输出；各传感器目录封装具体采集。"),
    ("p", "本机私有config.h提供Wi-Fi、MQTT、CA和SEND_INTERVAL_MS，加入.gitignore而不提交。协议常量与私密配置分离，使团队能够审查设备行为而不暴露账号。PlatformIO环境使用普通esp32dev和Arduino框架，协议ID中的_s3只为历史兼容，不代表开发板型号。"),
    ("h3", "4.3.2 主循环五阶段"),
    ("bullet", "阶段一检查Wi-Fi。断开时先通知MQTT和水枪执行统一安全停止，再尝试重连，并用板载LED提示网络状态。"),
    ("bullet", "阶段二调用mqtt_loop，处理连接事件、完整收包队列、JSON解析、时间和值域校验以及命令分发，不在网络回调中直接移动舵机。"),
    ("bullet", "阶段三推进WaterGunController，完成舵机稳定等待、定时截止和动态超时。"),
    ("bullet", "阶段四推进DHT11、JW01和保留的LED模块；BH1750连续模式在遥测时读取当前光照。"),
    ("bullet", "阶段五按SEND_INTERVAL_MS发布QoS 1遥测快照，使用独立时间基准，不因一次网络重连改变采样节奏。"),
    ("h3", "4.3.3 MQTT收包与幂等"),
    ("p", "ESP-IDF MQTT事件回调可能分片到达数据，因此驱动先按total_data_len、current_data_offset和data_len拼接完整JSON，再放入深度受限的FreeRTOS队列。主循环消费队列后使用ArduinoJson解析[6]，防止网络任务被舵机等待或复杂校验阻塞。"),
    ("p", "设备接受UUID v4或1至20位十进制command_id。最近命令缓存保存ID、FNV-1a内容指纹和最终ACK：同ID同内容不再操作GPIO，只重发原ACK；同ID不同内容拒绝。异步水枪命令进入Pending，待800毫秒稳定并实际决定泵输出后再写入最终缓存。"),
    ("h3", "4.3.4 水枪状态机实现"),
    ("image", "watergun"),
    ("caption", "图8 二维水枪安全执行状态机"),
    ("p", "mapBearingToPan对-90至90度进行线性插值，mapRangeToTilt对300至1200毫米进行反向线性插值。targetChanged使用0.1的浮点容差吸收JSON表示误差；位置不变时跳过panTilt写入，避免重复PWM造成舵机来回动作。"),
    ("p", "定时喷水先验证设备时间已经同步、spray_ends_at仍在未来且与duration误差合理，再同时设置Epoch和millis截止。动态喷水保存session_id、last_sequence和last_command_ms；超过3秒无合法新序号时只停泵并保留会话序号，用于继续拒绝延迟旧包。"),
    ("h2", "4.4 MQTT协议实现"),
    ("h3", "4.4.1 连接参数"),
    ("p", "设备和后端均使用MQTT 3.1.1、TLS 8883和QoS 1。设备Client ID为sab-dev-{device_id}，持久会话保留未确认消息；后端使用独立Client ID。设备配置LWT到status Topic，异常断开后Broker保留offline快照。ESP-MQTT提供事件驱动客户端、QoS与TLS能力[12]。"),
    ("h3", "4.4.2 命令验证顺序"),
    ("p", "设备按精确command Topic接收后，依次检查JSON语法、command_id、expires_at、operation、target和值类型。grow_light检查0至90整数；target_position检查方位、距离、模式、喷水设置、泵百分比、会话和定时字段的组合关系。任何校验失败都在不写GPIO的前提下返回稳定错误码。"),
    ("p", "主要错误包括INVALID_COMMAND、CAPABILITY_UNSUPPORTED、COMMAND_EXPIRED、DEVICE_TIME_UNSYNCED、TARGET_OUT_OF_RANGE、STALE_TARGET、TIMED_SPRAY_EXPIRED、COMMAND_QUEUE_FULL、ACTUATOR_INTERLOCK和INTERNAL_ERROR。后端将错误映射为HTTP状态和命令资源状态，Web展示具体原因。"),
    ("h3", "4.4.3 ACK语义"),
    ("p", "command_ack至少包含command_id、state、actual和error。executed表示ESP32已经把输出写入或状态机完成，无位置反馈、水流反馈和命中反馈，因此feedback_verified=false。若需要宣称真实角度或流量正确，必须增加角度、流量或压力传感器。"),
    ("h2", "4.5 后端实现"),
    ("h3", "4.5.1 FastAPI应用生命周期"),
    ("p", "backend_api/main.py创建FastAPI应用并注册知识库、农业来源、应用状态、病害图片、设备、站点、助手、摄像头和水枪路由。生命周期中初始化数据库、MQTT、监控、摄像头与外部服务，关闭时依次停止后台任务、断开Broker并释放摄像头。"),
    ("p", "依赖包括FastAPI、Uvicorn、SQLAlchemy、paho-mqtt、httpx、OpenCV Headless和Pillow。所有真实配置从.env与本机EMQX覆盖文件读取；MQTT_URI必须为mqtts且端口8883，旧MQTT_HOST、1883和本地Mosquitto不会静默回退。"),
    ("h3", "4.5.2 MQTT运行时"),
    ("p", "mqtt_service.py解析单一MQTT_URI，建立持久EMQX连接并订阅设备上行通配主题。收到消息后先解析Topic，再把遥测、状态、能力和ACK分派到领域服务。发布命令时固定QoS 1、retain=false，并记录发布是否被客户端接受。"),
    ("p", "后端不是简单消息透传。它校验设备能力、站点归属、命令范围与动作互锁，生成command_id和expires_at，持久化命令资源，等待ACK并向SSE发布command_update。这样浏览器和C5不需要接触EMQX账号，也不能绕过服务端安全策略。"),
    ("image", "backend"),
    ("caption", "图9 FastAPI后端消息、领域服务、数据库与终端链路"),
    ("h3", "4.5.3 站点、水枪与监控服务"),
    ("p", "site_service.py负责遥测标准化、站点快照、历史时间桶、设备状态、能力、命令与共享助手会话。water_gun_service.py维护静态和动态会话、sequence、倒计时与MQTT命令构造。monitoring_service.py维护报警阈值、在线状态和遥测接收质量。"),
    ("p", "动态模式中，页面heartbeat首先表示Web会话存活。目标未变化且未喷水时不发布MQTT，避免舵机重复动作；正在喷水时后端为设备生成新sequence作为安全保活。后端停止、会话过期或用户明确停止都会生成spray_enabled=false命令。"),
    ("h3", "4.5.4 知识、视觉与助手编排"),
    ("p", "kb_service把知识内容切分为chunk并执行本地检索，agri_source_service管理外部农业来源，vision_service和crop_vision_service处理图像分析，deepseek_service调用云端大模型。assistant_orchestrator把上下文、检索、工具和候选动作组织为一次可追踪的对话回合。"),
    ("p", "C5语音由speech_service处理ASR和TTS，音频响应带有自定义头部元数据。无论Web文本还是C5语音，涉及硬件的动作都进入EdgeAssistantAction记录，只有confirm决定后才调用设备或水枪服务；cancel、过期和设备离线均不会下发命令。"),
    ("image", "ai"),
    ("caption", "图10 知识库、视觉、大模型与用户确认闭环"),
    ("h2", "4.6 Web前端实现"),
    ("h3", "4.6.1 Vue状态和组件"),
    ("p", "src/main.ts挂载应用，App.vue维护八个业务视图、站点状态、水枪会话、助手线程、知识库、天气、报警和持久化设置。CameraGrowthPanel封装摄像头预览与抓帧，EChartPanel封装历史图表，MarkdownContent使用marked与DOMPurify渲染助手内容，避免不可信HTML直接进入页面。"),
    ("p", "api.ts中的requestJson统一设置基础URL、超时和错误翻译；上传、SSE、助手流和摄像头资源使用专门函数。默认站点为greenhouse_001，默认设备为greenhouse_001_s3。运行配置通过VITE_API_BASE_URL和VITE_SITE_ID注入。"),
    ("h3", "4.6.2 页面生命周期和资源释放"),
    ("p", "页面挂载后恢复仪表盘状态，加载站点、历史、报警、设备健康、天气、摄像头配置、知识库和水枪状态，再建立SSE。运行中维护5秒状态校准、60秒历史同步、水枪倒计时和按需AI分析。"),
    ("p", "组件卸载时关闭EventSource、清理interval和timeout、停止动态水枪heartbeat、释放Blob URL并结束浏览器语音识别。Web开发服务器和FastAPI应使用项目的一键关闭脚本终止整个进程树，避免后端子进程继续占用摄像头。"),
    ("h2", "4.7 项目目录与配置管理"),
    ("h3", "4.7.1 目录职责"),
    ("p", "仓库根目录下docs保存跨端标准、问题和比赛文档；esp32保存普通ESP32 PlatformIO工程；前端同时保存Vue应用、FastAPI后端、数据库、运行脚本和集成文档。C5源码应单独归档到esp32c5_voice_display，不与普通ESP32工程混放。"),
    ("p", "构建产物、虚拟环境、node_modules、数据库临时文件、私密config.h、EMQX证书与环境变量文件均由根.gitignore统一管理。README说明安装、启动、停止、摄像头释放和常见故障，避免各子目录继续生成重复的临时说明。"),
    ("h3", "4.7.2 开发与运行环境"),
    ("p", "普通ESP32使用VS Code PlatformIO和esp32dev环境；compile_commands.json用于C/C++跳转，但属于生成物。Web需要Node.js 20.19及以上，后端使用Python虚拟环境安装requirements.txt。SQLite随后端自动创建，不存在单独的数据库启动命令。"),
    ("p", "一键启动脚本应把后端和Web作为可跟踪后台进程启动，使终端仍可输入命令；一键关闭脚本按PID和端口检查并终止子进程，随后确认8000端口、Vite端口和摄像头句柄均已释放。"),
    ("h2", "4.8 测试与验证"),
    ("h3", "4.8.1 自动化和构建验证"),
    ("p", "普通ESP32最近一次PlatformIO构建已经完成编译和链接，记录中的RAM占用约16.3%、Flash占用约75.4%。该结论证明当前依赖、头文件和主固件语法可构建，不证明GPIO电平、舵机角度、水流和EMQX端到端链路已经正确。"),
    ("p", "后端最近一次完整pytest记录为165项通过并包含7个子测试，警告主要来自依赖弃用提示。测试覆盖设备API、MQTT Topic解析、命令发布与ACK、站点状态、水枪静态/动态/定时模式、助手动作确认、语音服务、位置与视觉服务。Vue前端也已经通过类型检查与Vite生产构建。"),
    ("h3", "4.8.2 真机验收方法"),
    ("bullet", "传感器验收：逐个接入DHT11、BH1750和JW01，对照独立仪表检查单位、有效范围、采样间隔、断线quality和5秒遥测。"),
    ("bullet", "执行器验收：断泵测试GPIO27与GPIO13方向和限位；测量GPIO14、GPIO26驱动电平与PWM；确认舵机独立供电和共地。"),
    ("bullet", "通信验收：观察EMQX五类Topic的QoS、retain与触发时机，逐条关联后端command_id、设备串口和command_ack。"),
    ("bullet", "故障验收：测试重复命令、同ID异内容、过期命令、旧sequence、Wi-Fi断开、后端退出、Web关闭、定时结束和动态超时均不会留下持续喷水。"),
    ("bullet", "产品验收：以实际植物位置标定方位和距离，记录不同泵功率的落点与流量，再替换UN_CALIBRATED_PLACEHOLDER常量。"),
    ("h2", "4.9 领域数据与接口实现细节"),
    ("h3", "4.9.1 站点状态对象"),
    ("p", "SiteState是Web和C5获取现场事实的核心对象。它聚合site_id、sampled_at、received_at、四类传感器、执行器状态、普通ESP32在线状态、能力、水枪会话和最近命令。sampled_at表示设备采样时刻，received_at表示后端接收时刻，两者不能混用；前者用于曲线与数据新鲜度，后者用于网络延迟和接收监控。"),
    ("p", "传感器对象由value、unit和quality组成。value允许数值或null，quality为invalid或unsupported时页面不应把null格式化为0。执行器状态表示设备最后上报的实际快照；如果用户刚提交命令但还没有ACK或新遥测，页面可显示“等待设备确认”，不能提前覆盖快照。"),
    ("p", "设备在线状态来自retained status与LWT，并结合后端最后接收时间判断。capabilities决定页面是否允许显示正式控制入口：普通ESP32声明DHT11温湿度、BH1750、JW01、grow_light和water_gun；未声明的执行器不会因为前端类型中存在旧字段就自动启用。"),
    ("h3", "4.9.2 命令资源生命周期"),
    ("p", "命令资源通常经历created、published、executed、rejected、expired或timed_out状态。created表示后端已完成请求校验并写入数据库；published表示MQTT客户端接受发布，不代表Broker送达；executed或rejected必须来自匹配command_id的设备ACK；timed_out表示约定时限内没有最终ACK。"),
    ("p", "HTTP 202用于表示命令已进入异步处理，而不是动作已经完成。Web收到202后订阅command_update或轮询状态，直到终态。若设备离线、能力不支持、参数越界或动作互锁，后端在发布前返回4xx；若MQTT运行时不可用或队列失败，返回服务错误并保持数据库中的失败原因。"),
    ("p", "后端记录命令来源，例如Web手动、助手确认或后端策略，但精简MQTT报文不重复该字段。这样ESP32只处理与安全执行有关的数据，审计和权限信息保留在可信服务端。一个逻辑动作只创建一个command_id，重试传输仍使用同一ID，防止用户刷新页面造成物理重复执行。"),
    ("h3", "4.9.3 水枪领域对象"),
    ("p", "WaterGunState保存mode、ground_range_mm、bearing_deg、spray_enabled、pump_control_percent、spray_schedule、remaining_seconds、session_id、sequence、target_source和stop_reason。该对象描述后端理解的会话状态，而ESP32 telemetry中的water_gun描述端侧实际输出，两者通过命令和ACK逐步收敛。"),
    ("p", "静态状态没有session_id，sequence为0；动态状态必须有session_id并按设备可见更新递增。target_source可以标记manual、camera或assistant，便于审计目标来源，但不会改变ESP32角度映射。stop_reason区分用户停止、定时完成、动态超时、网络故障和命令拒绝，帮助页面解释为什么水泵关闭。"),
    ("p", "preview接口只执行范围校验和目标计算，不写数据库命令、不发布MQTT。static接口用于固定目标；dynamic/start创建会话；dynamic/target更新位置；dynamic/spray只改变喷射状态；dynamic/heartbeat只维护页面会话；dynamic/stop结束会话。将这些意图拆开可以避免一个频繁指针事件同时切换喷水。"),
    ("h3", "4.9.4 数据库写入与去重"),
    ("p", "telemetry_receipts按message_id记录接收，用于丢弃QoS 1重发；site_history_samples按站点和时间桶保存用于曲线的标准化样本；site_snapshots保存每个站点最新完整状态。三类数据分别服务去重、历史和快速读取，不能只用一张无限增长的日志表承担所有用途。"),
    ("p", "device_commands和site_commands记录不同接口产生的命令与最终ACK，EdgeAssistantAction再关联助手候选动作。数据库事务先保存业务记录，再发布事件；若MQTT发布失败，错误状态同样落库，页面能够看到失败，而不是因为异常抛出丢失审计信息。"),
    ("p", "知识库使用knowledge_bases、knowledge_items和knowledge_chunks分离集合、原始条目和检索片段。更新条目时重新生成片段，删除知识库时清理从属数据。病害图片元数据和分析结果写入disease_photos，原始图片文件由受控目录管理，数据库不直接保存大体积视频流。"),
    ("h2", "4.10 异常处理与运行可观测性"),
    ("h3", "4.10.1 ESP32串口诊断"),
    ("p", "固件串口统一使用带模块和阶段的日志，例如BOOT、SENSOR、TELEMETRY、PUMP、LAMP、WATER_GUN和MQTT。启动日志打印物理芯片、协议设备ID、GPIO和初始化结果；传感器日志区分INIT_OK、READ_OK和READ_FAIL；执行器日志同时打印requested、actual、duty和reason。"),
    ("p", "水枪日志能够还原完整状态机：MAP_PLACEHOLDER记录输入位置和映射角，SERVO_WAIT记录等待，TARGET_UNCHANGED说明跳过PWM，TIMER_ARM说明本地截止，DYNAMIC_TIMEOUT说明安全停泵，PENDING_CANCEL说明旧异步命令被互锁取消。通过command_id可以与后端和EMQX记录逐条关联。"),
    ("p", "诊断日志不应打印Wi-Fi密码、MQTT密码、CA私钥或云端API Key。高频telemetry只在5秒调度点输出一次，避免串口日志本身阻塞主循环。真机问题排查先确认安全初始化和网络，再确认Topic收发，最后观察GPIO日志与物理动作。"),
    ("h3", "4.10.2 后端健康与错误隔离"),
    ("p", "后端健康检查分别观察HTTP服务、数据库、MQTT、摄像头和外部模型。某个外部农业来源、天气服务或大模型失败时，只影响对应功能并返回清晰错误，不应停止MQTT接收或水枪安全服务。MQTT断开时命令传输失败，但SQLite、历史查询和知识库仍可用。"),
    ("p", "摄像头状态接口返回设备索引、打开状态、分辨率、帧率和最近错误。若配置索引错误或设备被其他进程占用，Web显示摄像头不可用，而不是统一显示“云端连接中断”。后端关闭时记录release结果，关闭脚本再检查残留Python进程。"),
    ("p", "服务端日志记录请求路径、关键业务ID和异常阶段，但对环境变量和用户上传内容进行必要脱敏。助手链路分别标记检索、模型、ASR和TTS失败，使开发者能够判断是知识不足、模型凭据、网络还是音频格式问题。"),
    ("h3", "4.10.3 Web错误状态"),
    ("p", "Web把网络错误、HTTP业务拒绝、设备离线、ACK超时、摄像头不可用和AI服务失败作为不同状态展示。requestJson解析后端detail和稳定错误码，避免所有问题都变成“网络异常”。SSE断线只影响实时推送，5秒REST校准仍可恢复；REST也失败时才标记云端不可达。"),
    ("p", "控制按钮在请求进行中进入busy状态，防止重复提交；命令返回202后显示等待设备确认；收到executed ACK后显示输出已应用；收到rejected、expired或timed_out时恢复可操作并展示原因。页面卸载不会假装命令取消，真实停止必须调用后端接口。"),
    ("p", "摄像头MJPEG使用独立资源URL，不经过JSON请求。图片加载失败时组件读取camera/status获取具体原因，并允许重新配置索引。Blob URL、事件流和语音识别在组件卸载时释放，避免页面多次切换造成内存、麦克风或摄像头资源残留。"),

    ("h1", "5 其他内容"),
    ("h2", "5.1 安全与可靠性设计"),
    ("h3", "5.1.1 电气和机械安全"),
    ("p", "ESP32只输出逻辑信号，不直接驱动水泵、灯具和舵机负载。水泵应设置反向电动势保护、保险或限流，电源容量需覆盖两个SG90同时启动的峰值电流；水路与电路分区布置，接头处设置防漏和固定。"),
    ("p", "水平舵机0至180度和垂直舵机5至60度是软件上限，实际机构仍可能因安装偏心、连杆长度和喷嘴尺寸提前碰撞。机械端点标定必须停泵、低速、逐步逼近，并在最终装配状态重复验证。"),
    ("h3", "5.1.2 软件失效安全"),
    ("p", "设备采用默认关闭、停止优先、命令有效期、动态超时、本地定时、幂等和能力校验。网络回调不直接操作硬件，所有动作在主循环统一执行。即使后端或页面错误发送重复目标，ESP32也会按command_id、sequence和位置变化判断阻止重复动作。"),
    ("p", "安全并不只由设备端承担。Web对高风险喷水提供确认，后端检查设备在线和能力，EMQX使用最小ACL，设备执行后再由ACK更新页面。任何一层失败都应停止或拒绝动作，而不是为了“看起来在线”回退到模拟数据。"),
    ("h2", "5.2 部署与运行"),
    ("h3", "5.2.1 启动顺序"),
    ("p", "部署前配置后端虚拟环境、Web依赖、EMQX TLS账号、普通ESP32私密config.h和云端模型密钥。启动后端后先确认SQLite可写、EMQX已连接和摄像头索引可打开，再启动Web；ESP32上电后应依次看到安全初始化、Wi-Fi连接、MQTT连接、status、capabilities和周期telemetry。"),
    ("p", "C5上线时还需配置独立设备ID、语音服务参数和允许的Topic。后端和两个嵌入式端必须使用相同Topic前缀与时钟基准；命令有效期依赖SNTP，设备时间未同步时应保持拒绝控制而不是忽略时间校验。"),
    ("h3", "5.2.2 停止与恢复"),
    ("p", "停止系统应先通过Web或后端发送水枪停止，再运行一键关闭脚本终止FastAPI和Vite进程树，最后检查端口和摄像头占用。普通ESP32在MQTT或Wi-Fi断开时会本地停泵；重新连接后只接收未过期的新命令，不执行retained历史控制。"),
    ("p", "摄像头仍被占用时，应根据后端启动记录和端口定位实际Python进程，而不是只关闭浏览器。后端生命周期和关闭脚本必须调用release；若进程被强制中断，则由关闭脚本终止整个子进程树，防止Uvicorn重载子进程残留。"),
    ("h2", "5.3 标定、局限与实现边界"),
    ("h3", "5.3.1 当前需要实物标定的参数"),
    ("p", "水枪仍有四类关键标定：水平舵机0至180度与真实方向的对应、目标距离与垂直5至60度的关系、目标距离和水压与泵百分比的关系、舵机稳定等待时间。当前800毫秒、40%最小泵功率和线性距离模型用于保证软件流程可执行，不能作为最终产品精度指标。"),
    ("p", "标定应记录设备安装高度、喷嘴方向、供电电压、泵功率、目标坐标和实际落点，使用多组重复测量拟合曲线。若单段线性模型误差较大，可在不改变MQTT契约的情况下，把ESP32内部映射替换为分段线性或查表模型。"),
    ("h3", "5.3.2 当前不可直接宣称的能力"),
    ("p", "没有角度反馈传感器时，不能从ACK证明SG90真实到达角度；没有流量或压力传感器时，不能证明水泵产生目标流量；没有落点识别闭环时，不能证明水枪命中指定植物。本文把这些能力表述为控制输出已应用或待真机标定。"),
    ("p", "当前工作区没有C5固件源码，因此不能依据后端接口宣称屏幕和语音终端已经完成编译与真机验收。真实EMQX是否成功通信也需要连接日志、Broker消息和设备ACK共同证明，仅有后端发布代码或ESP32串口心跳不足以构成端到端证据。"),
    ("h2", "5.4 可维护性与扩展性"),
    ("h3", "5.4.1 契约演进"),
    ("p", "新增传感器时先定义领域字段、单位、quality和能力声明，再实现设备采集、后端持久化和前端展示；新增执行器时先确定安全状态、值域、ACK和互锁。不能因为页面已有旧控件就把风机、加热器等不存在设备重新加入正式能力。"),
    ("p", "协议版本通过Topic前缀演进。v1内保持字段语义兼容，新增可选字段必须有默认行为；破坏性变化应使用新版本Topic并提供迁移期。后端作为适配层可以同时订阅多个设备版本，但Web继续使用稳定REST领域模型。"),
    ("h3", "5.4.2 数据与服务扩展"),
    ("p", "单机SQLite满足当前家庭原型和竞赛演示，未来可迁移PostgreSQL并增加备份、用户鉴权和多站点隔离。摄像头可以从PC外接设备迁移到网络摄像头或边缘相机，只要继续由后端提供统一状态、帧和分析接口。"),
    ("p", "知识库可按植物种类、病害和生长阶段扩展，检索结果应保留来源和更新时间。大模型供应商可以替换，但助手编排、动作确认和设备契约保持独立，避免模型接口变化影响硬件控制。"),
    ("h2", "5.5 应用价值"),
    ("h3", "5.5.1 家庭养护价值"),
    ("p", "系统把连续环境观察、定点操作和专业知识转化为家庭用户可理解的状态与建议。用户不在家时仍可查看数据和画面，回家后也能通过现场屏幕或语音快速了解植物情况，减少因遗忘、误判或盲目浇水造成的养护失败。"),
    ("h3", "5.5.2 教学与工程价值"),
    ("p", "项目覆盖嵌入式驱动、实时状态机、MQTT、TLS、FastAPI、数据库、Vue、视觉、知识库和大模型编排，且以统一契约连接各端。它不仅展示功能，也体现了设备安全、幂等、数据质量、配置隔离和测试边界等完整物联网工程方法。"),
    ("h2", "5.6 后续优化"),
    ("h3", "5.6.1 近期优化方向"),
    ("bullet", "完成水枪落点、泵功率、舵机脉宽和稳定时间的实物标定，并把结果写入版本化标定文件。"),
    ("bullet", "将C5固件工程归档，补齐屏幕状态、语音端到端时延、断网恢复和最终ACK展示测试。"),
    ("bullet", "增加水流、压力或舵机位置反馈，使command_ack从输出已应用升级为部分物理闭环验证。"),
    ("bullet", "为真实EMQX配置最小ACL、证书轮换和连接监控，并保存端到端验收记录。"),
    ("bullet", "完善用户鉴权、站点隔离、数据库备份和操作审计，为多家庭设备部署做准备。"),
    ("h2", "5.7 项目协同与交付管理"),
    ("h3", "5.7.1 多端职责与交付物"),
    ("p", "普通ESP32交付物包括PlatformIO源码、统一契约常量、私密配置说明、串口诊断、构建记录和真机GPIO验收记录；C5交付物包括屏幕页面、音频链路、设备Topic、显示状态契约和语音真机测试；后端交付物包括数据库模型、MQTT网关、REST/SSE、知识和智能服务、环境变量模板与自动化测试；Web交付物包括生产构建、接口客户端、八个业务视图和异常状态说明。"),
    ("p", "各端不能只报告“代码已完成”。交付状态至少区分源码存在、依赖可安装、编译或构建通过、单元测试通过、与真实EMQX联通、与真实设备联通以及完整用户流程验收。跨端问题由command_id、message_id、session_id和时间戳关联，减少依赖截图或口头描述判断。"),
    ("h3", "5.7.2 契约优先的协同方法"),
    ("p", "MQTT Topic、JSON、REST、单位、错误码、GPIO和机械范围以smartagribrain-v1-standard.md为统一来源。开发者修改某一端前先判断是否改变契约：实现缺陷可以在本端修复；字段语义、单位和值域变化必须同步评审并更新各端测试，不能在前端或ESP32单方面增加兼容字段。"),
    ("p", "所有未确认问题集中记录在docs/open-questions.md，每个问题包含背景、影响、需要谁回答、可选结论和明确答复位置。收到回答后把结论迁移到正式契约和实现，问题条目标记已解决；不能长期留在“已收到答复待实现”而没有代码、测试或文档证据。"),
    ("p", "历史测试Topic、模拟数据、旧传感器和不存在执行器只保留在必要的迁移说明中，不进入正式运行路径。测试若需要替身，应限定在tests目录并通过依赖注入使用，生产配置禁止自动回退到本地Mosquitto、模拟遥测或虚假摄像头。"),
    ("h3", "5.7.3 版本、配置和发布检查"),
    ("p", "发布前固定普通ESP32固件版本、后端版本、Web构建版本和协议版本，并记录对应Git提交。固件capabilities携带版本，后端日志记录启动配置摘要，Web显示后端连接状态。发生问题时能够确定是代码不一致、配置不一致还是硬件标定不一致。"),
    ("p", "配置检查包括EMQX域名、8883端口、TLS CA、账号ACL、设备ID、站点ID、Topic前缀、系统时间、摄像头索引、数据库路径、模型密钥和语音服务参数。检查只输出是否存在和非敏感摘要，不回显密码。任何曾提交到Git或公开文档的密钥都应立即轮换。"),
    ("p", "最终演示按固定顺序执行：上电确认执行器关闭；观察status、capabilities和连续telemetry；打开Web检查实时与历史；执行低亮度补光并等待ACK；断泵移动两个舵机；执行短时低功率定点喷水；测试动态停止；展示摄像头、知识库和助手；最后运行关闭脚本并确认摄像头释放。每一步保留日志和结果。"),
    ("p", "验收记录应包含测试日期、固件与服务版本、硬件接线、EMQX环境、命令JSON、设备串口、后端日志、Web终态和实物观察。对于未通过项目，记录预期、实际、复现步骤和安全恢复方法；修复后使用相同用例回归。通过这种证据链，可以区分“报文正确但硬件接线错误”“设备执行但页面未更新”“页面提交但Broker未发布”等不同问题，也能防止后续修改重新引入舵机重复动作、摄像头残留占用或模拟路径回退。验收结果应由执行人复核并随版本归档，不能只保存在临时终端输出或聊天记录中。"),

    ("h1", "6 参考文献"),
]


REFERENCES = [
    ("OASIS. MQTT Version 3.1.1 Plus Errata 01[EB/OL]. 2015. https://docs.oasis-open.org/mqtt/mqtt/v3.1.1/errata01/os/mqtt-v3.1.1-errata01-os-complete.html.", "OASIS2015MQTT", "OASIS", "MQTT Version 3.1.1 Plus Errata 01", "2015", "https://docs.oasis-open.org/mqtt/mqtt/v3.1.1/errata01/os/mqtt-v3.1.1-errata01-os-complete.html"),
    ("ASAIR. DHT11 Temperature and Humidity Sensor Module[EB/OL]. 2026. https://www.aosong.com/en/Products/info.aspx?itemid=2257.", "ASAIRDHT11", "ASAIR", "DHT11 Temperature and Humidity Sensor Module", "2026", "https://www.aosong.com/en/Products/info.aspx?itemid=2257"),
    ("ROHM Semiconductor. BH1750FVI Ambient Light Sensor IC[EB/OL]. 2011. https://www.rohm.com/products/sensors-mems/ambient-light-sensors/bh1750fvi-product.", "ROHM2011BH1750", "ROHM Semiconductor", "BH1750FVI Ambient Light Sensor IC", "2011", "https://www.rohm.com/products/sensors-mems/ambient-light-sensors/bh1750fvi-product"),
    ("TowerPro. SG90 9g Micro Servo Datasheet[EB/OL]. 2026. https://components101.com/sites/default/files/component_datasheet/SG90%20Servo%20Motor%20Datasheet.pdf.", "TowerProSG90", "TowerPro", "SG90 9g Micro Servo Datasheet", "2026", "https://components101.com/sites/default/files/component_datasheet/SG90%20Servo%20Motor%20Datasheet.pdf"),
    ("Espressif Systems. Arduino-ESP32 LED Control API[EB/OL]. 2026. https://docs.espressif.com/projects/arduino-esp32/en/latest/api/ledc.html.", "Espressif2026LEDC", "Espressif Systems", "Arduino-ESP32 LED Control API", "2026", "https://docs.espressif.com/projects/arduino-esp32/en/latest/api/ledc.html"),
    ("Blanchon B. ArduinoJson Documentation[EB/OL]. 2026. https://arduinojson.org/.", "Blanchon2026ArduinoJson", "Benoit Blanchon", "ArduinoJson Documentation", "2026", "https://arduinojson.org/"),
    ("FastAPI. Server-Sent Events[EB/OL]. 2026. https://fastapi.tiangolo.com/tutorial/server-sent-events/.", "FastAPI2026SSE", "FastAPI", "Server-Sent Events", "2026", "https://fastapi.tiangolo.com/tutorial/server-sent-events/"),
    ("Vue.js. Vue 3 Guide: Introduction[EB/OL]. 2026. https://cn.vuejs.org/guide/introduction.", "Vue2026Guide", "Vue.js", "Vue 3 Guide: Introduction", "2026", "https://cn.vuejs.org/guide/introduction"),
    ("SQLite. SQLite Documentation[EB/OL]. 2026. https://www.sqlite.org/docs.html.", "SQLite2026Docs", "SQLite", "SQLite Documentation", "2026", "https://www.sqlite.org/docs.html"),
    ("EMQ Technologies. EMQX Documentation[EB/OL]. 2026. https://docs.emqx.com/en/emqx/latest/.", "EMQX2026Docs", "EMQ Technologies", "EMQX Documentation", "2026", "https://docs.emqx.com/en/emqx/latest/"),
    ("OpenCV. OpenCV Documentation[EB/OL]. 2026. https://docs.opencv.org/4.x/.", "OpenCV2026Docs", "OpenCV", "OpenCV Documentation", "2026", "https://docs.opencv.org/4.x/"),
    ("Espressif Systems. ESP-MQTT Programming Guide[EB/OL]. 2026. https://docs.espressif.com/projects/esp-idf/en/latest/esp32/api-reference/protocols/mqtt.html.", "Espressif2026MQTT", "Espressif Systems", "ESP-MQTT Programming Guide", "2026", "https://docs.espressif.com/projects/esp-idf/en/latest/esp32/api-reference/protocols/mqtt.html"),
]


def patch_bibliography(path: Path) -> None:
    tmp = path.with_suffix(".bib.tmp.docx")
    ns = {"b": BIB_NS}
    with ZipFile(path, "r") as source:
        root = etree.fromstring(source.read("customXml/item1.xml"))
        for child in list(root):
            root.remove(child)
        for _, tag, corporate, title, year, url in REFERENCES:
            src = etree.SubElement(root, f"{{{BIB_NS}}}Source")
            values = {
                "Tag": tag,
                "SourceType": "InternetSite",
                "Guid": "{" + str(uuid.uuid4()).upper() + "}",
            }
            for key, value in values.items():
                node = etree.SubElement(src, f"{{{BIB_NS}}}{key}")
                node.text = value
            author_outer = etree.SubElement(src, f"{{{BIB_NS}}}Author")
            author_inner = etree.SubElement(author_outer, f"{{{BIB_NS}}}Author")
            corp = etree.SubElement(author_inner, f"{{{BIB_NS}}}Corporate")
            corp.text = corporate
            for key, value in {
                "Title": title,
                "Year": year,
                "Publisher": corporate,
                "URL": url,
                "YearAccessed": "2026",
                "MonthAccessed": "7",
                "DayAccessed": "25",
            }.items():
                node = etree.SubElement(src, f"{{{BIB_NS}}}{key}")
                node.text = value
        bib_xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        with ZipFile(tmp, "w", compression=ZIP_DEFLATED) as target:
            for item in source.infolist():
                target.writestr(item, bib_xml if item.filename == "customXml/item1.xml" else source.read(item.filename))
    tmp.replace(path)


def set_picture_alt(paragraph, description: str) -> None:
    doc_pr = paragraph._p.find(".//wp:docPr", namespaces={"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"})
    if doc_pr is not None:
        doc_pr.set("descr", description)
        doc_pr.set("title", description)


def main() -> None:
    if len(sys.argv) not in {4, 5}:
        raise SystemExit("usage: expand_design_doc.py INPUT.docx BACKUP.docx ASSET_DIR [OUTPUT.docx]")
    if sys.argv[1] == "AUTO":
        candidates = [item for item in Path("docs").glob("*.docx") if not item.name.startswith("~$")]
        preferred = [item for item in candidates if item.name == "物联网设计.docx"]
        if len(preferred) == 1:
            candidates = preferred
        if len(candidates) != 1:
            raise RuntimeError(f"Expected one docx in docs, found {len(candidates)}")
        path = candidates[0].resolve()
    else:
        path = Path(sys.argv[1]).resolve()
    output_path = Path(sys.argv[4]).resolve() if len(sys.argv) == 5 else path
    backup = Path(sys.argv[2]).resolve()
    asset_dir = Path(sys.argv[3]).resolve()
    backup.parent.mkdir(parents=True, exist_ok=True)
    asset_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(path, backup)

    images = generate_refined_figures(asset_dir)
    for name, image_path in images.items():
        if not image_path.exists():
            raise RuntimeError(f"Missing figure asset {name}: {image_path}")

    document = Document(path)
    if len(document.paragraphs) < 15:
        raise RuntimeError("Unexpected template structure")
    replace_paragraph_text(document.paragraphs[12],
        "面向家庭阳台、室内种植空间和小型庭院中环境观察不连续、养护知识分散及浇灌位置难以准确控制等问题，"
        "本文设计并实现SmartAgriBrain智慧农业系统。产品采用方形一体化“小型智慧种植屋”结构，将DHT11温湿度、"
        "BH1750光照、JW01二氧化碳传感器，补光模块、二维舵机云台、定点灌溉水枪、ESP32-C5语音显示终端和普通ESP32"
        "采集执行节点集成于同一设备。普通ESP32通过Wi-Fi、TLS和云端EMQX以MQTT QoS 1持续上报环境与执行器状态，"
        "并以有效期、幂等、ACK和非阻塞安全状态机执行补光及定点喷水。FastAPI后端负责消息接入、SQLite持久化、REST/SSE、"
        "报警、摄像头、知识库、视觉分析和大模型编排；Vue Web提供实时监测、历史查询、设备控制、病害分析与知识管理；"
        "C5终端用于现场显示和云端语音问答。系统以统一领域数据和通信契约连接各端，形成集环境监测、精准执行、可视化展示、"
        "远程管理、语音交互和智能决策于一体的家庭智能植物管家。")
    replace_paragraph_text(document.paragraphs[13], "关键词：家庭智能种植；ESP32；MQTT；二维云台；精准灌溉；知识库；语音交互")
    remove_after(document, 14)

    heading_items = [(kind, text) for kind, text in CONTENT if kind in {"h1", "h2", "h3"}]
    bookmark_names = [f"toc_{index:03d}" for index in range(1, len(heading_items) + 1)]
    for (kind, text), bookmark in zip(heading_items, bookmark_names):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt({"h1": 0, "h2": 18, "h3": 36}[kind])
        paragraph.paragraph_format.space_after = Pt(1.5)
        add_hyperlink(paragraph, text, bookmark)
        for run in paragraph.runs:
            set_run_font(run, east_asia="宋体", size=10)

    heading_cursor = 0
    chapter_count = 0
    for kind, value in CONTENT:
        if kind in {"h1", "h2", "h3"}:
            level = {"h1": 1, "h2": 2, "h3": 3}[kind]
            paragraph = document.add_heading(value, level=level)
            bookmark_paragraph(document, paragraph, bookmark_names[heading_cursor])
            heading_cursor += 1
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
            if level == 1:
                chapter_count += 1
                paragraph.paragraph_format.page_break_before = True
            continue
        if kind == "p":
            add_body(document, value)
        elif kind == "bullet":
            add_bullet(document, value)
        elif kind == "code":
            add_code(document, value)
        elif kind == "image":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(images[value]), width=Inches(6.35))
            set_picture_alt(paragraph, {
                "product": "SmartAgriBrain小型智慧种植屋产品概念图，展示一体化种植空间、补光、传感器、二维水枪、储水与显示模块。",
                "modules": "系统模块协同概念图，展示植物终端、普通ESP32、C5、传感器、摄像头与云端服务之间的联系。",
                "coordinates": "喷水位置与双舵机角度映射图，左侧为方位负90至正90度到水平舵机0至180度，右侧为距离300至1200毫米到垂直舵机60至5度。",
                "mqtt": "EMQX五类Topic图，展示telemetry、status、capabilities、command和command_ack的方向、触发方式与retain策略。",
                "architecture": "SmartAgriBrain端到端科研架构图，展示现场层、边缘层、消息服务层、数据智能层和用户交互层。",
                "message": "遥测与控制双向时序图，展示普通ESP32、EMQX、FastAPI数据库和Web或C5之间的七步消息流。",
                "gpio": "普通ESP32硬件GPIO图，展示DHT11、BH1750、JW01、补光灯、水泵和双SG90的引脚及PWM参数。",
                "watergun": "二维水枪安全状态机图，展示接收、校验、停泵、位置映射、双轴转向、稳定等待、建立保护和开泵确认顺序。",
                "backend": "FastAPI后端数据处理图，展示MQTT接入、协议校验、领域服务、SQLite、REST SSE和Vue C5链路。",
                "ai": "知识增强智能流程图，展示实时环境、历史、视觉与知识进入助手编排、云端模型和用户确认门。",
            }[value])
        elif kind == "caption":
            paragraph = document.add_paragraph(value)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(8)
            for run in paragraph.runs:
                set_run_font(run, east_asia="宋体", size=9)

    # Reference heading is the final h1 in CONTENT. Add references immediately after it.
    for index, (display, *_rest) in enumerate(REFERENCES, 1):
        paragraph = document.add_paragraph(f"[{index}] {display}")
        paragraph.paragraph_format.left_indent = Pt(21)
        paragraph.paragraph_format.first_line_indent = Pt(-21)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.25
        for run in paragraph.runs:
            set_run_font(run, size=9.5)
        bookmark_paragraph(document, paragraph, f"ref_{index}")

    document.core_properties.title = "SmartAgriBrain智慧农业系统设计说明书"
    document.core_properties.subject = "全国大学生物联网设计竞赛设计说明书"
    document.core_properties.comments = "扩展版：依据当前普通ESP32、FastAPI、Vue、EMQX与统一v1契约整理。"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    patch_bibliography(output_path)

    print(f"saved={output_path}")
    print(f"backup={backup}")
    print(f"headings={len(heading_items)}")
    print(f"references={len(REFERENCES)}")
    print(f"bytes={output_path.stat().st_size}")


if __name__ == "__main__":
    main()
